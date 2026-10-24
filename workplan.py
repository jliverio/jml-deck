"""
workplan.py - draws the Phase Architect v2 plan as native PowerPoint slides
on the template's "Default" layout:

  1. Timeline by workstream  (Gantt rows grouped by phase + governance row)
  2. Deliverables by workstream (what each stream hands over, due week, where reviewed)
  3. Steering meetings (one column per check-in / workshop / gate, max 4 per slide)

Input = the `plan` object returned by validate_plan_v2.js. All weeks are
already computed there; this module only draws. Titles are composed here
from the numbers (no model writes them), so they can never disagree with the chart.

    add_workplan_slides(prs, plan, section="Approach", client="[Client]") -> [slides]
    build_workplan_deck(plan) -> pptx bytes (template with only these slides; for tests)
"""
import io
import os

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt, Inches
from lxml import etree

NAVY = RGBColor(0x05, 0x1C, 0x2C)
BLUE = RGBColor(0x22, 0xA3, 0xDF)
ELECTRIC = RGBColor(0x22, 0x22, 0xF6)
TEAL = RGBColor(0x18, 0x77, 0xA2)
LIGHT = RGBColor(0xA2, 0xD5, 0xEC)
YELLOW = RGBColor(0xF8, 0xBC, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
LGREY = RGBColor(0xD0, 0xD0, 0xD0)
PHASE_COLORS = [BLUE, ELECTRIC, TEAL, NAVY, GREY]  # a fifth phase gets its own colour (Oct 18: five-phase plans)
FONT = "Arial"

L = Emu(554736)            # content left (title placeholder left)
W = Emu(11082528)          # content width
TOP = Emu(1700000)         # content top (below subtitle)
BOTTOM = Emu(6300000)      # content bottom (above footer)


# ------------------------------------------------------------------ helpers
def _box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, dash=False, line_w=0.75):
    s = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # sysDash
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Emu(45720); tf.margin_top = tf.margin_bottom = Emu(18288)
    return s


def _text(shape_or_slide, text, size=10, bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
          x=None, y=None, w=None, h=None, italic=False, wrap=True):
    if x is not None:
        sh = shape_or_slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        sh.text_frame.margin_left = sh.text_frame.margin_right = Emu(0)
        sh.text_frame.margin_top = sh.text_frame.margin_bottom = Emu(0)
    else:
        sh = shape_or_slide
    tf = sh.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = str(ln)
        f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = FONT; f.color.rgb = color
    return sh


def _bullets(shape, items, size=11, color=BLACK, space_after=6):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "171450"); pPr.set("indent", "-171450")
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•")
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        r = p.add_run(); r.text = str(it)
        f = r.font; f.size = Pt(size); f.name = FONT; f.color.rgb = color


def _line(slide, x1, y1, x2, y2, color=BLACK, width=0.75, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1)), Emu(int(y1)), Emu(int(x2)), Emu(int(y2)))
    c.line.color.rgb = color; c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    return c


def _new_slide(prs, title, subtitle, section, slide=None, after=None):
    """Use the template's marked slide when given; otherwise add a slide on the
    "Default" layout (after `after` when given). The tracker (idx 17) is only
    written on new slides: on template slides the renderer has filled it."""
    if slide is not None:
        s = slide
        for ph in s.placeholders:
            idx = ph.placeholder_format.idx
            if idx == 0:
                ph.text_frame.text = title
            elif idx == 1:
                ph.text_frame.text = subtitle
        return s
    layout = after.slide_layout if after is not None else prs.slide_layouts[0]
    s = prs.slides.add_slide(layout)
    if after is not None:
        lst = prs.slides._sldIdLst
        ids = list(lst)
        pos = [i for i, sl in enumerate(prs.slides) if sl.slide_id == after.slide_id][0]
        el = ids[-1]
        lst.remove(el)
        lst.insert(pos + 1, el)
        for ph in after.placeholders:
            if ph.placeholder_format.idx == 17:
                section = ph.text_frame.text or section
    for ph in s.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text_frame.text = title
        elif idx == 1:
            ph.text_frame.text = subtitle
        elif idx == 17:
            ph.text_frame.text = section
    return s


# Number style for every code-written text (José, Oct 18): counts up to nine in words, weeks in digits,
# ranges with an en dash ("two phases", "12 workstreams", "12–16 weeks"). The Summary Writer's check uses the same.
NUM_WORDS = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine"]


def _num(n):
    return NUM_WORDS[n] if 0 <= n < len(NUM_WORDS) else str(n)


def _count(n, noun, plural=None):
    return f"{_num(n)} {noun if n == 1 else (plural or noun + 's')}"


def _weeks(dw):
    mn, mx = dw["min"], dw["max"]
    return (f"{mx} week" + ("" if mx == 1 else "s")) if mn == mx else f"{mn}–{mx} weeks"


def _cap(s):
    return s[:1].upper() + s[1:]


_weeks_label = _weeks


def _phase_color(order):
    return PHASE_COLORS[(order - 1) % len(PHASE_COLORS)]


def _clip(s, n):
    s = " ".join(str(s or "").split())
    return s if len(s) <= n else s[: n - 1].rstrip(" ,.;:") + "…"


# ------------------------------------------------------------------ titles
def _client_name(plan):
    c = str(plan.get("client") or "").strip()
    return c if c and not c.startswith("[") else "The client"


def _gate_weeks(plan):
    """[(phase, gate week)] in phase order; the gate sits in the phase's last week."""
    gates = {e["phase_order"]: e for e in plan["governance"]["events"] if e["type"] == "gate"}
    phases = sorted(plan["phases"], key=lambda p: p["order"])
    return [(p, (gates.get(p["order"]) or {}).get("week") or p.get("end_week")) for p in phases]


def _week_list(ws):
    ws = [str(w) for w in ws]
    return ("week " + ws[0]) if len(ws) == 1 else "weeks " + ", ".join(ws[:-1]) + " and " + ws[-1]


def roadmap_title(plan, limit=110):
    """When the client decides (approach overview): the gates, not the duration, which slide 4 gives."""
    gw = _gate_weeks(plan)
    client = _client_name(plan)
    last_p, last_w = gw[-1]
    if len(gw) == 1:
        label = (last_p.get("gate") or {}).get("label") or "final decision"
        return f"{client} takes one decision, at the {label} in week {last_w}"
    head = (f"in week {gw[0][1]} whether {gw[1][0]['name']} goes ahead" if len(gw) == 2
            else f"in {_week_list([w for _, w in gw[:-1]])} whether the next phase goes ahead")
    t = f"{client} decides {head}; the work closes in week {last_w}"
    if len(t) <= limit:
        return t
    nxt = gw[1][0]["name"] if len(gw) == 2 else "Each next phase"
    return f"{nxt} goes ahead only after a decision in {_week_list([w for _, w in gw[:-1]])}; the work closes in week {last_w}"


def _peak_parallel(plan):
    wss = plan.get("workstreams") or []
    total = plan["total_duration_weeks"]["max"]
    return max([sum(1 for w in wss if w["start_week"] <= wk <= w["end_week"]) for wk in range(1, total + 1)] or [0])


def _max_steering_gap(plan):
    wk = sorted({e["week"] for e in plan["governance"]["events"]})
    return max((b - a for a, b in zip(wk, wk[1:])), default=None)


def titles(plan):
    """Slide titles composed from the plan's numbers, so a title cannot disagree with its chart. The template's
    title variables (approach_title_01, ...) carry them; a value given for a variable replaces the composed text
    (plan["_titles"], set by the renderer)."""
    nws = len({str(w.get("name", "")).strip().lower() for w in plan["workstreams"]}) or len(plan["workstreams"])
    nd = len(plan["deliverables"])
    ev = [e for e in plan["governance"]["events"] if e["type"] != "kickoff"]
    gates = sum(1 for e in ev if e["type"] == "gate")
    peak = _peak_parallel(plan)
    gap = _max_steering_gap(plan)
    run = "The workstreams run one at a time" if peak <= 1 else f"Up to {_count(peak, 'workstream')} run in the same week"
    ws_part = "The workstream hands over" if nws == 1 else f"The {_count(nws, 'workstream')} hand over"
    if len(ev) == 1:
        gov = "After the kick-off, one steering meeting reviews the work and closes it with a decision"
    else:
        closing = "each closing a phase with a decision" if gates == len(ev) else (
            f"{_num(gates)} of them closing a phase with a decision")
        gov = f"After the kick-off, {_count(len(ev), 'steering meeting')} review the work, {closing}"
    out = {
        "roadmap": roadmap_title(plan),
        "timeline": run + (f", with no more than {gap} weeks between steering meetings" if gap else ""),
        "deliverables": f"{ws_part} {_count(nd, 'deliverable')}, each reviewed at a steering meeting",
        "governance": gov,
    }
    for k, v in (plan.get("_titles") or {}).items():
        if k in out and str(v or "").strip():
            out[k] = str(v).strip()
    return out


# ------------------------------------------------------------------ layout rules
# Everything below scales with the plan: the week axis with the duration, the
# rows with the number of workstreams, and both slides split into "continued"
# pages instead of shrinking text below a readable size.
MIN_ROW = Inches(0.3)        # timeline row: smallest height that still holds a one-line label
MAX_ROW = Inches(0.7)
MIN_LINE = Inches(0.34)      # deliverables slide: one deliverable line (2-line stream label fits)
MAX_LINE = Inches(0.46)


def week_step(total):
    """Weeks per labelled tick: about 16 labels at most."""
    for st in (1, 2, 4, 8, 13, 26):
        if total / st <= 16:
            return st
    return 52


def _paginate(groups, cap):
    """Pack groups (lists) into the fewest pages of at most `cap` items, as even
    as possible (no lone straggler page). A group stays together unless it alone
    is larger than a page."""
    import math
    groups = [list(g) for g in groups if g]
    total = sum(len(g) for g in groups)
    if total <= cap:
        return [[x for g in groups for x in g]] if total else [[]]

    def pack(limit):
        pages, cur = [], []
        for g in groups:
            if cur and len(cur) + len(g) > limit:
                pages.append(cur); cur = []
            while len(g) > limit:
                pages.append(g[:limit]); g = g[limit:]
            cur = cur + g
        if cur:
            pages.append(cur)
        return pages

    n = math.ceil(total / cap)
    while True:
        for limit in range(math.ceil(total / n), cap + 1):
            pages = pack(limit)
            if len(pages) <= n:
                return pages
        n += 1


def _ws_order(plan):
    return sorted(plan["workstreams"], key=lambda w: (w["phase_order"], w["start_week"], w["id"]))


# timeline geometry (fixed parts)
_BAND_H = Inches(0.24)
_HEAD_H = Inches(0.26)
_GRID_TOP = TOP + _BAND_H + Inches(0.04) + _HEAD_H
_STEER_GAP = Inches(0.32)          # grid bottom -> steering markers
_LABEL_TIER = Inches(0.42)
_LABEL_H = Inches(0.4)
_LEGEND_H = Inches(0.25)
_DEADLINE_H = Inches(0.3)


def _timeline_capacity(plan):
    reserve = _STEER_GAP + Inches(0.16) + _LABEL_TIER + _LABEL_H + _LEGEND_H
    if (plan.get("deadline") or {}).get("weeks_available") is not None:
        reserve += _DEADLINE_H
    avail = BOTTOM - reserve - _GRID_TOP
    return avail, max(1, int(avail // MIN_ROW))


# ------------------------------------------------------------------ time units (template 1.5)
# José, 24 Oct: the grid is 26 cm wide and holds at most 24 boxes (about 1.09 cm each). One box is one unit and
# every box is the same unit: weeks while the plan fits 24 of them, then months, bi-months, quarters, four-month
# periods and semesters. Workstreams, phases, deliverables and meetings snap to whole units: in a first draft it
# does not matter whether a workstream takes 6 weeks or 2 months. Past 24 semesters there is no unit (the S3
# layout then leaves the plan slides out).
MAX_BOXES = 24
MAX_WEEKS = MAX_BOXES * 26                   # 24 semesters (12 years): past this the S3 deck has no plan slides
UNITS = (  # units per year, header, singular, plural
    (52, "Week", "week", "weeks"), (12, "Month", "month", "months"), (6, "Bi-month", "bi-month", "bi-months"),
    (4, "Quarter", "quarter", "quarters"), (3, "Four-month", "four-month period", "four-month periods"),
    (2, "Semester", "semester", "semesters"))


def _rnd(x):
    return int(x + 0.5)                       # round half up, so week 6 of a 52-week year is month 1 (1.38)


def _box_end(week, k):
    return max(1, _rnd(week * k / 52))


def _span(s, e, k):
    a = 1 if s <= 1 else _rnd((s - 1) * k / 52) + 1
    return a, max(a, _rnd(e * k / 52))


def _to_units(plan, k):
    """The plan with every week turned into a box number at k units per year (k=52: unchanged)."""
    import copy
    g = copy.deepcopy(plan)
    prev = 0
    pspan = {}
    for p in sorted(g["phases"], key=lambda p: p["order"]):
        a, b = _span(p["start_week"], p["end_week"], k)
        a = max(a, prev + 1)
        b = max(a, b)
        p["start_week"], p["end_week"] = a, b
        pspan[p["order"]] = (a, b)
        prev = b
    wspan = {}
    for w in g["workstreams"]:
        ps, pe = pspan.get(w["phase_order"], (1, prev))
        a, b = _span(w["start_week"], w["end_week"], k)
        a = min(max(a, ps), pe)
        b = min(max(b, a), pe)
        w["start_week"], w["end_week"] = a, b
        wspan[w["id"]] = (a, b)
    for d in g["deliverables"]:
        if d.get("due_week"):
            a, b = wspan.get(d.get("workstream_id"), (1, prev))
            d["due_week"] = min(max(_box_end(d["due_week"], k), a), b)
    # one marker per box: a decision gate wins, then a workshop, the kick-off, a check-in
    rank = {"gate": 0, "workshop": 1, "kickoff": 2, "checkin": 3}
    best = {}
    for e in g["governance"]["events"]:
        if e["type"] == "gate" and e.get("phase_order") in pspan:
            e["week"] = pspan[e["phase_order"]][1]
        else:
            e["week"] = min(_box_end(e["week"], k), prev)
        cur = best.get(e["week"])
        if cur is None or rank.get(e["type"], 9) < rank.get(cur["type"], 9):
            best[e["week"]] = e
    g["governance"]["events"] = sorted(best.values(), key=lambda e: e["week"])
    taken = set(best)
    g["governance"]["status_weeks"] = sorted({min(_box_end(wk, k), prev) for wk in g["governance"].get("status_weeks") or []}
                                             - taken)
    g["total_duration_weeks"] = {"min": prev, "max": prev}
    return g


def time_unit(plan):
    """(units per year, header, singular, plural, the plan in those units), or None past 24 semesters."""
    total = plan["total_duration_weeks"]["max"]
    for k, head, one, many in UNITS:
        if _span(1, total, k)[1] > MAX_BOXES:
            continue
        g = _to_units(plan, k)
        if g["total_duration_weeks"]["max"] <= MAX_BOXES:
            return k, head, one, many, g
    return None


# ------------------------------------------------------------------ slide 1: timeline
def timeline_slides(prs, plan, section, client, slide=None, unit_scale=False):
    """One or more timeline pages (each holds whole phases where possible). unit_scale: template 1.5's
    rule, at most 24 boxes of one unit (time_unit); otherwise one column per week, as in 1.4."""
    unit = time_unit(plan) if unit_scale else None
    avail, cap = _timeline_capacity(plan)
    order = _ws_order(plan)
    numbering = {w["id"]: i + 1 for i, w in enumerate(order)}
    groups = [[w for w in order if w["phase_order"] == p["order"]] for p in plan["phases"]]
    pages = _paginate([g for g in groups if g], cap)
    out = []
    for i, rows in enumerate(pages):
        out.append(timeline_slide(prs, plan, section, client, rows=rows, numbering=numbering, avail=avail,
                                  page=(i, len(pages)), slide=slide if i == 0 else None,
                                  after=(out[-1] if (i and slide is not None) else None), unit=unit))
    return out


def timeline_slide(prs, plan, section, client, slide=None, rows=None, numbering=None, avail=None,
                   page=(0, 1), after=None, unit=None):
    T = titles(plan)
    ranged = any(p["duration_weeks"]["min"] != p["duration_weeks"]["max"] for p in plan["phases"])
    title = T["timeline"] + (" (continued)" if page[0] else "")
    head, one, many = ("Week", "week", "weeks") if unit is None else unit[1:4]
    sub = f"Plan by workstream and {one}" + (" (drawn on the longer estimate)" if ranged else "")
    if page[1] > 1:
        sub += f" – page {page[0] + 1} of {page[1]}"
    s = _new_slide(prs, title, sub, section, slide=slide, after=after)
    grid = plan if unit is None else unit[4]          # positions in boxes; titles and the deadline stay in weeks
    total = grid["total_duration_weeks"]["max"]
    phases = grid["phases"]
    events = grid["governance"]["events"]
    wss = rows if rows is not None else _ws_order(plan)
    if unit is not None:
        gws = {w["id"]: w for w in grid["workstreams"]}
        wss = [gws.get(w["id"], w) for w in wss]
    if numbering is None:
        numbering = {w["id"]: i + 1 for i, w in enumerate(_ws_order(plan))}
    if avail is None:
        avail = _timeline_capacity(plan)[0]

    _text(s, f"Initial view – to confirm with {client}", size=10, italic=True, color=GREY,
          align=PP_ALIGN.RIGHT, x=L + W - Inches(4), y=TOP - Inches(0.28), w=Inches(4), h=Inches(0.22))

    label_w = Inches(1.85)
    gx = L + label_w
    gw = W - label_w
    col = gw / total
    step = week_step(total) if unit is None else 1
    n_rows = len(wss)
    row_h = max(MIN_ROW, min(MAX_ROW, avail / max(n_rows, 1)))
    y0 = TOP
    # phase band
    for p in phases:
        x = gx + (p["start_week"] - 1) * col
        w = (p["end_week"] - p["start_week"] + 1) * col
        b = _box(s, x, y0, w, _BAND_H, fill=_phase_color(p["order"]))
        est = "" if p["duration_weeks"]["min"] == p["duration_weeks"]["max"] else " (estimate)"
        if unit is None or unit[0] == 52:
            full = f"{p['name']} · {_weeks_label(p['duration_weeks'])}{est}"
        else:
            nb = p["end_week"] - p["start_week"] + 1
            full = f"{p['name']} · {nb} {one if nb == 1 else many}"
        fit = int((w - Inches(0.2)) / Inches(0.075))  # inside the box margins (Oct 18: 'Measurement · 6 weeks' overflowed a 6-of-39-week band)
        _text(b, _clip(full if len(full) <= fit else p["name"], fit), size=10 if fit >= 12 else 8,
              bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _text(s, "Phase", size=10, bold=True, color=NAVY, x=L, y=y0, w=label_w, h=_BAND_H)
    # week header: one label per tick (every `step` weeks)
    yh = y0 + _BAND_H + Inches(0.04)
    hdr = _box(s, L, yh, W, _HEAD_H, fill=NAVY)
    _text(hdr, head, size=10, bold=True, color=WHITE)
    lab_size = 10 if col * step >= Inches(0.3) else 9
    for wk in range(1, total + 1, step):
        _text(s, str(wk), size=lab_size, bold=True, color=WHITE, align=PP_ALIGN.CENTER if step == 1 else PP_ALIGN.LEFT,
              x=gx + (wk - 1) * col + (0 if step == 1 else Emu(18288)), y=yh,
              w=col * min(step, total - wk + 1) if step == 1 else max(col * min(step, total - wk + 1) - Emu(18288), col),  # last tick stays inside the grid (Oct 18)
              h=_HEAD_H)
    yr = yh + _HEAD_H
    grid_top = yr
    grid_bottom = yr + n_rows * row_h
    # gridlines on ticks only, phase boundaries always
    bounds = {p["end_week"] for p in phases}
    for wk in range(1, total):
        is_boundary = wk in bounds
        if not is_boundary and wk % step:
            continue
        x = gx + wk * col
        _line(s, x, grid_top, x, grid_bottom, color=BLACK if is_boundary else LGREY,
              width=1.0 if is_boundary else 0.5, dash=not is_boundary)
    _line(s, gx, grid_top, gx, grid_bottom, color=BLACK, width=1.0)
    my = grid_bottom + _STEER_GAP
    d = max(Inches(0.1), min(Inches(0.2), col * 0.8))
    for e in events:
        if e["type"] in ("workshop", "gate"):
            cx = gx + (e["week"] - 0.5) * col
            _line(s, cx, grid_top, cx, my - d / 2, color=GREY, width=0.75, dash=True)

    # rows
    del_due = {dd["id"]: dd["due_week"] for dd in grid["deliverables"]}
    # stream labels always get two lines: 10pt in tall rows, 8pt in the tightest ones
    lsize = 10 if row_h >= Inches(0.4) else 8
    two_lines = True
    chars_line = int((label_w - Inches(0.12)) / Inches(0.083 if lsize == 10 else 0.067))
    for i, w in enumerate(wss):
        y = yr + i * row_h
        c = _phase_color(w["phase_order"])
        lab = _box(s, L, y + Emu(9144), label_w - Emu(27432), row_h - Emu(18288), fill=c)
        _text(lab, _clip(f"{numbering[w['id']]:02d} {w['name']}", chars_line * (2 if two_lines else 1)),
              size=lsize, bold=True, color=WHITE)
        bx = gx + (w["start_week"] - 1) * col + Emu(18288)
        bw = max(Emu(36576), (w["end_week"] - w["start_week"] + 1) * col - Emu(36576))
        _box(s, bx, y + row_h * 0.22, bw, row_h * 0.56, fill=c)
        # deliverable markers: a white diamond in the due week, its number inside.
        # Diamonds keep a readable size on long plans and are pushed apart instead of overlapping.
        dm = min(row_h * 0.72, Inches(0.3), max(col * 0.9, Inches(0.22)))
        marks = sorted(((del_due[did], did) for did in w.get("deliverable_ids", []) if did in del_due))
        placed = []
        for dw, did in marks:
            cx = gx + (dw - 0.5) * col
            if placed and cx < placed[-1] + dm * 0.95:
                cx = placed[-1] + dm * 0.95
            placed.append(cx)
        right = gx + gw - dm / 2 - Emu(9144)
        if placed and placed[-1] > right:          # shift the whole cluster back inside the grid
            shift = placed[-1] - right
            placed = [max(gx + dm / 2, c - shift) for c in placed]
        for (dw, did), cx in zip(marks, placed):
            dmd = _box(s, cx - dm / 2, y + (row_h - dm) / 2, dm, dm, fill=WHITE, line=NAVY, shape=MSO_SHAPE.DIAMOND, line_w=0.75)
            tf = dmd.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
            num = did[1:] if did[:1] == "D" else did
            _text(dmd, num, size=7 if len(num) > 1 else 8, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _line(s, L, y + row_h, L + W, y + row_h, color=LGREY, width=0.5, dash=True)
    # governance row
    _text(s, "Steering", size=10, bold=True, color=NAVY, x=L, y=my - d / 2, w=label_w, h=d)
    for wk in grid["governance"]["status_weeks"]:
        _box(s, gx + (wk - 0.5) * col - d / 2, my - d / 2, d, d, fill=WHITE, line=BLACK, shape=MSO_SHAPE.OVAL, line_w=0.75)
    gap = Emu(45720)
    evs = sorted(events, key=lambda e: e["week"])

    def place(lab_w):
        # two tiers of labels; when both are taken, the label goes on the freer tier, pushed right.
        # Returns the positions and whether any label still overlaps its neighbour at the right edge.
        last_right = [-10**12, -10**12]
        out, clash = [], False
        for e in evs:
            cx = gx + (e["week"] - 0.5) * col
            lx = min(max(cx - lab_w / 2, L), L + W - lab_w)
            free = [t for t in (0, 1) if lx > last_right[t] + gap]
            if free:
                tier = free[0]
            else:
                tier = 0 if last_right[0] <= last_right[1] else 1
                lx = min(last_right[tier] + gap, L + W - lab_w)
                clash = clash or lx < last_right[tier] + gap
            last_right[tier] = lx + lab_w
            out.append((e, lx, tier))
        return out, clash

    # crowded steering rows (many meetings, long plans) get narrower two-line labels (Oct 18: a 39-week,
    # five-phase plan put its last two gate labels on top of each other)
    for lab_w in (Inches(1.45), Inches(1.25), Inches(1.05), Inches(0.9)):
        placed, clash = place(lab_w)
        if not clash:
            break
    chars = int(2 * (lab_w - Inches(0.1)) / Inches(0.068))
    for e, lx, tier in placed:
        cx = gx + (e["week"] - 0.5) * col
        fill = YELLOW if e["type"] in ("workshop", "gate") else NAVY
        _box(s, cx - d / 2, my - d / 2, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
        ly = my + d / 2 + Inches(0.06) + tier * _LABEL_TIER
        _text(s, _clip(e["label"], min(34, chars)), size=9, bold=e["type"] == "gate", color=BLACK, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.TOP, x=lx, y=ly, w=lab_w, h=_LABEL_H)
    dl = plan.get("deadline")
    if dl and dl.get("weeks_available") is not None:
        t = plan["total_duration_weeks"]["max"]
        fits = "fits within" if t <= dl["weeks_available"] else "does NOT fit within"
        _text(s, f"Deadline: {dl.get('text') or 'stated in the brief'}. The plan ({t} weeks) {fits} the {dl['weeks_available']} weeks available, allowing two weeks to contract and start.",
              size=9, bold=True, color=NAVY, x=L, y=BOTTOM - _LEGEND_H - _DEADLINE_H + Inches(0.05), w=W, h=Inches(0.2))
    # legend
    ly = BOTTOM - Inches(0.18)
    lx = L
    cad = plan['governance']['status_cadence_weeks']
    for fill, line, name in ((NAVY, None, "Kick-off / check-in"), (YELLOW, None, "Steering workshop / decision gate"),
                             (WHITE, BLACK, ("Status meeting (weekly)" if cad == 1 else f"Status meeting (every {cad} weeks)"))):
        _box(s, lx, ly, Inches(0.13), Inches(0.13), fill=fill, line=line, shape=MSO_SHAPE.OVAL)
        _text(s, name, size=8, color=GREY, x=lx + Inches(0.18), y=ly - Inches(0.02), w=Inches(2.3), h=Inches(0.17))
        lx += Inches(2.55)
    _box(s, lx, ly, Inches(0.13), Inches(0.13), fill=WHITE, line=NAVY, shape=MSO_SHAPE.DIAMOND)
    _text(s, "Deliverable due (3 = D3, listed on the deliverables page)", size=8, color=GREY, x=lx + Inches(0.18), y=ly - Inches(0.02), w=Inches(3.7), h=Inches(0.17))
    return s


# ------------------------------------------------------------------ slide 2: deliverables
_DEL_COLS = (("Workstream", Inches(2.3)), ("Deliverables", Inches(5.25)), ("Due", Inches(0.9)), ("Reviewed at", None))
# What the client receives (José, Oct 19): an icon on each deliverable line, the words in a legend at the foot of the
# page. Same ids, labels and defaults as phase_architect/deliverable_formats.json (tests/test_formats.py checks it).
FORMATS = (("presentation", "Presentation"), ("report", "Report"), ("model", "Model"), ("dashboard", "Dashboard"),
           ("playbook", "Playbook"), ("training", "Training"), ("solution", "Working solution"))
FORMAT_DEFAULT = {"analysis": "presentation", "design": "presentation", "build": "solution",
                  "decision_pack": "presentation", "enablement": "training"}
ICON_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "format_icons")
_ICON = Inches(0.22)


def _format(d):
    f = d.get("format")
    return f if f in dict(FORMATS) else FORMAT_DEFAULT.get(d.get("kind"), "presentation")


def _icon(slide, fmt, x, y, size):
    path = os.path.join(ICON_DIR, fmt + ".png")
    if os.path.exists(path):
        pic = slide.shapes.add_picture(path, x, y, size, size)
        pic.name = f"format:{fmt}"
        return pic
    return None


def _format_legend(slide, plan):
    """Icons with their words, at the foot of the page, for the formats this plan uses (same on every page)."""
    used = {_format(d) for d in plan["deliverables"]}
    ly = BOTTOM - Inches(0.2)
    _text(slide, "What the client receives:", size=8, bold=True, color=GREY, x=L, y=ly, w=Inches(1.6), h=Inches(0.17))
    lx = L + Inches(1.6)
    for fid, label in FORMATS:
        if fid not in used:
            continue
        _icon(slide, fid, lx, ly, Inches(0.17))
        _text(slide, label, size=8, color=GREY, x=lx + Inches(0.21), y=ly - Inches(0.01), w=Inches(1.25), h=Inches(0.17))
        lx += Inches(0.3) + Inches(0.065) * len(label) + Inches(0.25)


def deliverables_slides(prs, plan, section, client, slide=None):
    """One or more deliverables pages; a workstream never splits across pages
    unless it alone has more deliverables than a page holds."""
    y_first = TOP - Inches(0.05) + Inches(0.38)
    avail = BOTTOM - _LEGEND_H - Inches(0.05) - y_first        # the format legend sits at the foot (Oct 19)
    cap = max(1, int(avail // MIN_LINE))
    order = _ws_order(plan)
    numbering = {w["id"]: i + 1 for i, w in enumerate(order)}
    lines = []                                     # (workstream, [deliverable ids on this line-group])
    groups = [[(w, did) for did in (w["deliverable_ids"] or [None])] for w in order]
    pages = _paginate(groups, cap)
    out = []
    for i, items in enumerate(pages):
        out.append(deliverables_slide(prs, plan, section, client, items=items, numbering=numbering, avail=avail,
                                      page=(i, len(pages)), slide=slide if i == 0 else None,
                                      after=(out[-1] if (i and slide is not None) else None)))
    return out


def deliverables_slide(prs, plan, section, client, slide=None, items=None, numbering=None, avail=None,
                       page=(0, 1), after=None):
    T = titles(plan)
    title = T["deliverables"] + (" (continued)" if page[0] else "")
    sub = "Deliverables by workstream" + (f" – page {page[0] + 1} of {page[1]}" if page[1] > 1 else "")
    s = _new_slide(prs, title, sub, section, slide=slide, after=after)
    phases = {p["order"]: p for p in plan["phases"]}
    dels = {d["id"]: d for d in plan["deliverables"]}
    order = _ws_order(plan)
    if numbering is None:
        numbering = {w["id"]: i + 1 for i, w in enumerate(order)}
    if items is None:
        items = [(w, did) for w in order for did in (w["deliverable_ids"] or [None])]
    cols = [(n, w if w is not None else W - Inches(8.45)) for n, w in _DEL_COLS]
    x = L
    y = TOP - Inches(0.05)
    for name, w in cols:
        _text(s, name, size=12, bold=True, color=BLACK, x=x, y=y, w=w, h=Inches(0.25))
        _line(s, x, y + Inches(0.3), x + w - Inches(0.08), y + Inches(0.3), color=BLACK, width=0.75)
        x += w
    y += Inches(0.38)
    if avail is None:
        avail = BOTTOM - _LEGEND_H - Inches(0.05) - y
    line_h = max(MIN_LINE, min(MAX_LINE, avail / max(len(items), 1)))
    size = 11 if line_h >= Inches(0.4) else 10
    icon_w = _ICON + Inches(0.1)
    name_chars = int((cols[1][1] - Inches(0.1) - icon_w) / Inches(0.0795 if size == 11 else 0.072))
    rev_chars = int((cols[3][1] - Inches(0.05)) / Inches(0.0795 if size == 11 else 0.072))
    # group consecutive items of the same workstream
    runs = []
    for w, did in items:
        if runs and runs[-1][0]["id"] == w["id"]:
            runs[-1][1].append(did)
        else:
            runs.append((w, [did]))
    for w, ids in runs:
        ids = [i for i in ids if i is not None]
        h = line_h * max(1, len(ids))
        c = _phase_color(w["phase_order"])
        lab = _box(s, L, y + Emu(18288), cols[0][1] - Emu(91440), h - Emu(36576), fill=c)
        _text(lab, [_clip(f"{numbering[w['id']]:02d} {w['name']}", 29), phases[w["phase_order"]]["name"] + " phase"], size=9, bold=True, color=WHITE)
        for j, did in enumerate(ids):
            d = dels[did]
            yy = y + j * line_h
            x = L + cols[0][1]
            _icon(s, _format(d), x, yy + (line_h - _ICON) / 2, _ICON)
            _text(s, _clip(f"{did}  {d['name']}", name_chars), size=size, x=x + icon_w, y=yy, w=cols[1][1] - Inches(0.1) - icon_w, h=line_h)
            x += cols[1][1]
            _text(s, d.get("due_label") or f"Week {d['due_week']}", size=size, x=x, y=yy, w=cols[2][1], h=line_h)
            x += cols[2][1]
            _text(s, _clip(d.get("presented_at") or "[Firm to complete]", rev_chars), size=size, x=x, y=yy, w=cols[3][1], h=line_h)
        y += h
        _line(s, L, y, L + W, y, color=LGREY, width=0.5, dash=True)
    _format_legend(s, plan)
    return s


# ------------------------------------------------------------------ slide 3: steering meetings
def governance_slides(prs, plan, section, client, per_slide=4, slide=None):
    T = titles(plan)
    dels = {d["id"]: d for d in plan["deliverables"]}
    evs = sorted([e for e in plan["governance"]["events"] if e["type"] != "kickoff"], key=lambda e: e["week"])
    # as few slides as `per_slide` columns allow, filled evenly with the extra meetings on the first slides
    # (5 meetings: 3 then 2; José, Oct 18), whatever phase each meeting closes; the subtitle names the phases
    pages = max(1, -(-len(evs) // per_slide))
    base, extra = divmod(len(evs), pages)
    chunks, start = [], 0
    for k in range(pages):
        size = base + (1 if k < extra else 0)
        chunks.append(evs[start:start + size]); start += size
    chunks = [c for c in chunks if c] or [[]]
    slides = []
    phase_by = {p["order"]: p for p in plan["phases"]}
    numbering = {}
    wk = 0
    for e in evs:
        if e["type"] == "workshop":
            wk += 1; numbering[e["id"]] = wk
    for ci, chunk in enumerate(chunks):
        ph_names = sorted({phase_by[e["phase_order"]]["name"] for e in chunk}, key=lambda n: [p["name"] for p in plan["phases"]].index(n))
        sub = "Steering meetings – " + " and ".join(ph_names) + (" phase" if len(ph_names) == 1 else " phases")
        title = T["governance"] if ci == 0 else T["governance"] + " (continued)"
        if slide is not None:
            s = _new_slide(prs, title, sub, section, slide=slide if ci == 0 else None,
                           after=None if ci == 0 else slides[-1])
        else:
            s = _new_slide(prs, title, sub, section)
        n = max(len(chunk), 3)
        gap = Inches(0.35)
        cw = (W - gap * (n - 1)) / n
        for i, e in enumerate(chunk):
            x = L + i * (cw + gap)
            y = TOP
            head = f"Week {e['week']}"
            lab_w = cw - Inches(0.7)
            fit11 = int(lab_w / Inches(0.088))
            lsize = 11 if len(e["label"]) <= fit11 else 9
            _text(s, _clip(e["label"], int(lab_w / Inches(0.072)) if lsize == 9 else fit11), size=lsize, bold=True,
                  x=x, y=y, w=lab_w, h=Inches(0.3))
            _text(s, head, size=10, color=GREY, align=PP_ALIGN.RIGHT, x=x + cw - Inches(0.68), y=y, w=Inches(0.68), h=Inches(0.3))
            _line(s, x, y + Inches(0.34), x + cw, y + Inches(0.34), color=BLACK, width=1.0)
            fill = YELLOW if e["type"] == "gate" else (ELECTRIC if e["type"] == "workshop" else NAVY)
            band = _box(s, x, y + Inches(0.42), cw, Inches(0.58), fill=fill)
            _text(band, e["purpose"], size=12, bold=True, color=NAVY if e["type"] == "gate" else WHITE)
            att = e.get("attendees") or []
            if att:
                _text(s, _clip("With: " + ", ".join(att), int((cw - Inches(0.2)) / Inches(0.066))), size=9, italic=True, color=GREY,
                      anchor=MSO_ANCHOR.TOP, x=x, y=y + Inches(1.06), w=cw, h=Inches(0.2))
            body = _box(s, x, y + Inches(1.3), cw, Inches(1.55))
            # long agendas (3 items of up to 80 characters in a narrow column) step down to 10 or 9 pt (Oct 18)
            asize = 11
            for asize in (11, 10, 9):
                cpl = max(10, int((cw - Inches(0.35)) / Inches(0.0075 * asize)))
                lines = sum(-(-len(str(a)) // cpl) for a in e["agenda"])
                if lines * Inches(0.0165 * asize) + len(e["agenda"]) * Pt(6) <= Inches(1.5):
                    break
            _bullets(body, e["agenda"], size=asize)
            kind_rank = {"decision_pack": 0, "design": 1, "build": 2, "analysis": 3, "enablement": 4}
            ordered = sorted([d for d in e["presents"] if d in dels], key=lambda d: (kind_rank.get(dels[d]["kind"], 9), -dels[d]["due_week"]))
            pres = [f"{d}  {dels[d]['name']}" for d in ordered]
            if pres:
                maxc = int((cw - Inches(0.2)) / Inches(0.066))  # one 9 pt line inside the box margins (Oct 18 stress test)
                shown = pres[:3] if len(pres) <= 3 else pres[:2]
                lines = [_clip(p, maxc) for p in shown] + ([f"+{len(pres) - len(shown)} more (see deliverables page)"] if len(pres) > len(shown) else [])
                _text(s, "Reviews", size=9, bold=True, color=GREY, anchor=MSO_ANCHOR.TOP, x=x, y=y + Inches(2.9), w=cw, h=Inches(0.2))
                tb = _text(s, lines, size=9, color=GREY, anchor=MSO_ANCHOR.TOP, x=x, y=y + Inches(3.1), w=cw, h=Inches(0.75))
            if e.get("decision"):
                dbox = _box(s, x, BOTTOM - Inches(0.62), cw, Inches(0.55), fill=None, line=NAVY)
                _text(dbox, ["Decision: " + e["decision"]], size=9, bold=True, color=NAVY)
            if i < len(chunk) - 1:
                _line(s, x + cw + gap / 2, y + Inches(0.4), x + cw + gap / 2, BOTTOM - Inches(0.05), color=GREY, width=0.75, dash=True)
        slides.append(s)
    return slides


# ------------------------------------------------------------------ slide 0: approach overview (template 1.4)


def _list(items):
    items = [str(i) for i in items if str(i).strip()]
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1] if items else ""


def roadmap_slide(prs, plan, section, client, slide=None):
    """Approach overview: one column per phase (what it answers, when it ends,
    the decision that closes it) and a governance strip. Everything is read
    from the plan; nothing here is written by a model."""
    s = _new_slide(prs, titles(plan)["roadmap"], "Approach overview: phases, decisions and how we work together",
                   section, slide=slide)
    phases = sorted(plan["phases"], key=lambda p: p["order"])
    events = plan["governance"]["events"]
    n = len(phases)
    gap = Inches(0.12)
    cw = (W - gap * (n - 1)) / n
    chev_h = Inches(0.62)
    y = TOP
    gates = {e["phase_order"]: e for e in events if e["type"] == "gate"}
    kick = next((e for e in events if e["type"] == "kickoff"), None)
    for i, p in enumerate(phases):
        x = L + i * (cw + gap)
        shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        ch = _box(s, x, y, cw + (gap if i < n - 1 else 0), chev_h, fill=_phase_color(p["order"]), shape=shape)
        ch.adjustments[0] = 0.25
        weeks = (f"Week {p['start_week']}" if p.get("start_week") == p.get("end_week")
                 else f"Weeks {p.get('start_week')}–{p.get('end_week')}") if p.get("start_week") else ""
        tf = ch.text_frame
        tf.margin_left = Inches(0.3) if i else Inches(0.15)
        _text(ch, [f"Phase {p['order']}: {p['name']}", f"{weeks} ({_weeks_label(p['duration_weeks'])})" if weeks
                   else _weeks_label(p["duration_weeks"])], size=12, bold=True, color=WHITE)
        tf.paragraphs[1].runs[0].font.bold = False
        tf.paragraphs[1].runs[0].font.size = Pt(10)
        # what it waits for (José, Oct 18: the sequence belongs on this slide)
        if i == 0:
            starts = f"Starts at the kick-off, week {kick['week'] if kick else 1}"
            if len(starts) > int((cw - Inches(0.1)) / Inches(0.075)):
                starts = f"Week {kick['week'] if kick else 1}, at the kick-off"
        else:
            pg = gates.get(phases[i - 1]["order"]) or {}
            plabel = (phases[i - 1].get("gate") or {}).get("label") or pg.get("label") or "previous decision"
            starts = f"Starts after the {plabel}, week {pg.get('week') or phases[i - 1].get('end_week')}"
            if len(starts) > int((cw - Inches(0.1)) / Inches(0.075)):   # narrow columns (5 phases): keep the week
                starts = f"Week {pg.get('week') or phases[i - 1].get('end_week')}, after the {plabel}"
            if len(starts) > int((cw - Inches(0.1)) / Inches(0.075)):   # still too long: name the phase, not the gate (Oct 19)
                starts = f"Week {pg.get('week') or phases[i - 1].get('end_week')}, after Phase {phases[i - 1]['order']}"
        _text(s, _clip(starts, int((cw - Inches(0.1)) / Inches(0.075))), size=10, italic=True, color=GREY,
              anchor=MSO_ANCHOR.TOP, x=x, y=y + chev_h + Inches(0.08), w=cw - Inches(0.1), h=Inches(0.22))
        # what the phase answers
        _text(s, "What it answers", size=9, bold=True, color=GREY, anchor=MSO_ANCHOR.TOP,
              x=x, y=y + chev_h + Inches(0.36), w=cw, h=Inches(0.2))
        _text(s, _clip(p.get("objective_one_line", ""), 170), size=12 if n < 5 else 11, anchor=MSO_ANCHOR.TOP,  # 5 columns: 11 pt (Oct 18)
              x=x, y=y + chev_h + Inches(0.58), w=cw - Inches(0.1), h=Inches(1.0))
        # the decision that closes it
        gy = y + chev_h + Inches(1.6)
        g = gates.get(p["order"]) or {}
        gate = p.get("gate") or {}
        label = gate.get("label") or g.get("label") or "Phase decision"
        decision = gate.get("decision") or g.get("decision") or ""
        wk = g.get("week") or p.get("end_week")
        _line(s, x, gy, x + cw, gy, color=BLACK, width=1.0)
        dm = _box(s, x, gy + Inches(0.14), Inches(0.24), Inches(0.24), fill=YELLOW, shape=MSO_SHAPE.DIAMOND)
        gfit = int((cw - Inches(0.34)) / Inches(0.085))
        full = f"{label}" + (f", week {wk}" if wk else "")
        short = len(full) > gfit and wk          # narrow columns (5 phases): the week moves to the decision line
        _text(s, _clip(label if short else full, gfit),
              size=11, bold=True, x=x + Inches(0.32), y=gy + Inches(0.1), w=cw - Inches(0.32), h=Inches(0.3))
        if decision:
            _text(s, _clip((f"Week {wk} decision: " if short else "Decision: ") + decision, 190), size=11 if n < 5 else 10, color=NAVY, anchor=MSO_ANCHOR.TOP,
                  x=x, y=gy + Inches(0.48), w=cw - Inches(0.1), h=Inches(1.0))
    # governance strip
    gov = plan["governance"]
    steer = [e for e in events if e["type"] != "kickoff"]
    who = []
    for e in events:
        for a in e.get("attendees") or []:
            if a and not str(a).startswith("[") and a not in who:
                who.append(a)
    cad = gov.get("status_cadence_weeks")
    first_due = min((d.get("due_week") for d in plan.get("deliverables") or [] if d.get("due_week")), default=None)
    cells = [
        ("Steering", _cap(_count(len(steer), "steering meeting")) + (f" with {_list(who)}" if who else "")
         + ("; each phase ends in a decision" if n > 1 else "; the work ends in a decision")),
        ("Working rhythm", (f"Kick-off in week {kick['week']}" if kick else "Kick-off at the start")
         + (f"; first deliverable in week {first_due}" if first_due else "")
         + (f"; status meetings every {cad} weeks" if cad and cad > 1 else ("; weekly status meetings" if cad == 1 else ""))),
        (f"What we ask of {client}", "A named project lead, plus time from the steering members; the exact commitment to agree at kick-off"),
    ]
    sy = BOTTOM - Inches(1.05)
    band = _box(s, L, sy, W, Inches(1.0), fill=None, line=NAVY)
    cwid = W / 3
    for k, (head, body) in enumerate(cells):
        cx = L + k * cwid
        _text(s, head, size=10, bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP,
              x=cx + Inches(0.15), y=sy + Inches(0.1), w=cwid - Inches(0.3), h=Inches(0.22))
        _text(s, _clip(body, 150), size=10, anchor=MSO_ANCHOR.TOP,
              x=cx + Inches(0.15), y=sy + Inches(0.36), w=cwid - Inches(0.3), h=Inches(0.6))
        if k:
            _line(s, cx, sy + Inches(0.12), cx, sy + Inches(0.88), color=LGREY, width=0.75)
    _text(s, f"Initial view – to confirm with {client}", size=10, italic=True, color=GREY,
          align=PP_ALIGN.RIGHT, x=L + W - Inches(4), y=TOP - Inches(0.28), w=Inches(4), h=Inches(0.22))
    return s


# ------------------------------------------------------------------ entry points
def add_workplan_slides(prs, plan, section="Approach", client="[Client]"):
    out = timeline_slides(prs, plan, section, client) + deliverables_slides(prs, plan, section, client)
    out += governance_slides(prs, plan, section, client)
    return out


MARKERS = ("roadmap", "timeline", "deliverables", "steering")


def marked_slides(prs):
    """{kind: slide} for template slides whose shape is named workplan:<kind>."""
    out = {}
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.name.startswith("workplan:") and sh.name[9:] in MARKERS:
                out[sh.name[9:]] = sl
    return out


def fill_template_slides(prs, plan, marked, client="[Client]", unit_scale=False):
    """Draw the plan onto the template's work-plan slides. Without a plan the
    slides stay and say what the firm has to complete."""
    if not (isinstance(plan, dict) and plan.get("phases")):
        todo = {"roadmap": "Approach overview: phases and decisions",
                "timeline": "Work plan: timeline by workstream", "deliverables": "Deliverables by workstream",
                "steering": "Steering meetings"}
        for kind, sl in marked.items():
            _new_slide(prs, "[Firm to complete] " + todo[kind], "", "", slide=sl)
        return list(marked.values())
    out = []
    if "roadmap" in marked:
        out.append(roadmap_slide(prs, plan, "", client, slide=marked["roadmap"]))
    if "timeline" in marked:
        out += timeline_slides(prs, plan, "", client, slide=marked["timeline"], unit_scale=unit_scale)
    if "deliverables" in marked:
        out += deliverables_slides(prs, plan, "", client, slide=marked["deliverables"])
    if "steering" in marked:
        out += governance_slides(prs, plan, "", client, slide=marked["steering"])
    return out


def build_workplan_deck(plan, client="[Client]", template=None):
    import varfill
    prs = varfill._load_template(template or varfill.TEMPLATE)
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
    add_workplan_slides(prs, plan, client=client)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    import json, sys
    data = json.load(open(sys.argv[1]))
    plan = data.get("plan", data)
    client = data.get("client", "[Client]")
    dest = sys.argv[2] if len(sys.argv) > 2 else "workplan.pptx"
    open(dest, "wb").write(build_workplan_deck(plan, client=client))
    print("wrote", dest)
