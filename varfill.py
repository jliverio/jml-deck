"""
varfill.py - JML AI Analyst deck renderer (variable-based, template 1.1).

The template marks every editable text with a lowercase variable name
(section_01, finding_02_headline, ...). Everything else ships as written.

    build_deck(values, rules=None) -> (pptx_bytes, warnings)

values  {variable: text}
rules   rows of the n8n `template_variables` table; each needs `variable`,
        `max_chars` and `if_empty`. Without rules the built-in defaults apply
        (offline use only).

Rules applied
  * only exact, whole variable names are replaced (title_01 never matches
    inside subtitle_01);
  * a variable split across text runs is re-joined when the pieces render
    alike; otherwise TemplateFormatError;
  * every variable in the template must have a table row (TemplateRulesError);
    a row with no variable in the template is only a warning;
  * if_empty: remove_row deletes the grouped row the variable sits in,
    hide_line deletes the line, clear blanks the whole line (label included),
    firm_todo prints [Firm to complete];
  * text longer than max_chars is cut at a word (or, for "; " lists, at a
    whole entry) and gets an ellipsis - a last resort; the workflow is meant
    to deliver text that already fits;
  * more agenda topics than the template holds -> TemplateCapacityError;
  * anything still unfilled at the end becomes [Firm to complete].
"""
import io
import os
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "template_v1_1.potx")
FIRM_TODO = "[Firm to complete]"
TOKEN = re.compile(r"(?<![A-Za-z0-9_])[a-z0-9]+(?:_[a-z0-9]+)+(?![A-Za-z0-9_])")
POLICIES = {"remove_row", "hide_line", "clear", "firm_todo"}

# Offline defaults (the live values come from the n8n table).
DEFAULT_RULES = [
    (r"cover_title_\d+", 21, "firm_todo"),
    (r"cover_subtitle_\d+", 100, "firm_todo"),
    (r"section_\d+", 20, "remove_row"),
    (r"objective_\d+", 127, "firm_todo"),
    (r"solution_\d+", 95, "firm_todo"),
    (r"outcome_\d+", 127, "firm_todo"),
    (r"finding_\d+_headline", 80, "remove_row"),
    (r"finding_\d+_action_taken", 90, "remove_row"),
    (r"finding_\d+_(context|result_achieved)", 90, "hide_line"),
    (r"sources_market_\d+", 204, "clear"),
    (r"next_step_\d+", 115, "remove_row"),
]


class TemplateCapacityError(ValueError):
    pass


class TemplateFormatError(ValueError):
    pass


class TemplateRulesError(ValueError):
    pass


# --------------------------------------------------------------- helpers
def _load_template(path):
    """python-pptx refuses .potx; patch the content type in memory."""
    with open(path, "rb") as fh:
        raw = fh.read()
    if path.endswith(".potx"):
        src = zipfile.ZipFile(io.BytesIO(raw))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml",
                                        b"presentationml.presentation.main+xml")
                dst.writestr(item, data)
        raw = buf.getvalue()
    return Presentation(io.BytesIO(raw))


def _shorten(s, n):
    s = " ".join(str(s).split())
    if not n or len(s) <= n:
        return s
    if "; " in s:  # a list: keep whole entries only
        out = ""
        for part in s.split("; "):
            nxt = part if not out else out + "; " + part
            if len(nxt) + 3 > n:
                break
            out = nxt
        if out:
            return out + "; …"
    cut = s[: n - 1]
    sp = cut.rfind(" ")
    if sp > n * 0.6:
        cut = cut[:sp]
    return cut.rstrip(" ,.;:") + "…"


def _fmt(run):
    f = run.font
    color = None
    try:
        if f.color and f.color.type is not None:
            color = str(f.color.rgb) if f.color.type == 1 else str(f.color.theme_color)
    except Exception:
        pass
    return (f.size, f.bold, f.italic, f.name, color)


def _same_look(runs):
    """True when runs render alike. Unset (inherited) matches anything;
    italic unset counts as not italic."""
    looks = [_fmt(r) for r in runs]
    for k in range(5):
        vals = {lk[k] for lk in looks}
        if k == 2:
            vals = {bool(v) for v in vals}
        else:
            vals.discard(None)
        if len(vals) > 1:
            return False
    return True


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _text_shapes(slide):
    return [s for s in _walk(slide.shapes) if getattr(s, "has_text_frame", False)]


def _join_split_tokens(para, where):
    """Re-join variables that PowerPoint split across runs."""
    for _ in range(50):
        runs = list(para.runs)
        bounds, pos = [], 0
        for r in runs:
            bounds.append((pos, pos + len(r.text)))
            pos += len(r.text)
        full = "".join(r.text for r in runs)
        split = None
        for m in TOKEN.finditer(full):
            idx = [i for i, (s, e) in enumerate(bounds) if e > s and s < m.end() and e > m.start()]
            if len(idx) > 1:
                split = (m.group(), idx)
                break
        if split is None:
            return
        name, idx = split
        if not _same_look([runs[i] for i in idx]):
            raise TemplateFormatError(
                f"{where}: variable '{name}' is split across runs with different formatting; "
                f"retype it in PowerPoint in one go")
        runs[idx[0]].text = "".join(runs[i].text for i in idx)
        for i in idx[1:]:
            runs[i].text = ""


def _remove(el):
    el.getparent().remove(el)


def _rule_table(rules):
    """{variable: (max_chars, if_empty)} from table rows, or None."""
    if not rules:
        return None
    table = {}
    for row in rules:
        name = str(row.get("variable") or "").strip()
        if not name:
            continue
        try:
            mx = int(row.get("max_chars")) if row.get("max_chars") not in (None, "") else None
        except (TypeError, ValueError):
            mx = None
        pol = str(row.get("if_empty") or "firm_todo").strip()
        table[name] = (mx, pol if pol in POLICIES else "firm_todo")
    return table


def _default_rule(name):
    for pat, mx, pol in DEFAULT_RULES:
        if re.fullmatch(pat, name):
            return mx, pol
    return None, "firm_todo"


# --------------------------------------------------------------- builder
def build_deck(values, rules=None, template=TEMPLATE):
    values = dict(values or {})
    warnings = []
    prs = _load_template(template)
    slides = list(prs.slides)
    table = _rule_table(rules)

    # 0. normalise split variables, collect what the template holds
    present = set()
    for si, sl in enumerate(slides, 1):
        for sh in _text_shapes(sl):
            for p in sh.text_frame.paragraphs:
                _join_split_tokens(p, f"slide {si} / {sh.name}")
                for r in p.runs:
                    present.update(TOKEN.findall(r.text))

    # 1. template and table must agree
    if table is not None:
        missing = sorted(v for v in present if v not in table)
        if missing:
            raise TemplateRulesError(
                "These template variables have no row in template_variables: " + ", ".join(missing))
        extra = sorted(v for v in table if v not in present)
        if extra:
            warnings.append(f"{len(extra)} table rows have no variable in the template "
                            f"(e.g. {extra[0]})")

    def rule(name):
        if table is not None and name in table:
            return table[name]
        return _default_rule(name)

    # 2. agenda capacity
    slots = {v for v in present if re.fullmatch(r"section_\d+", v)}
    asked = {k for k, v in values.items()
             if re.fullmatch(r"section_\d+", k) and str(v or "").strip()}
    if asked - slots:
        raise TemplateCapacityError(
            f"This proposal has {len(slots | asked)} agenda topics; the template holds {len(slots)}.")

    def val(name):
        v = values.get(name)
        v = "" if v is None else str(v).strip()
        if not v:
            return ""
        mx, _ = rule(name)
        out = _shorten(v, mx)
        if out != " ".join(v.split()):
            warnings.append(f"{name} cut to {mx} characters")
        return out

    cache = {}

    def value(name):
        if name not in cache:
            cache[name] = val(name)
        return cache[name]

    # 3. remove grouped rows: any empty remove_row variable, or all empty
    def prune(shapes):
        for sh in list(shapes):
            if sh.shape_type != MSO_SHAPE_TYPE.GROUP:
                continue
            names = set()
            for s in _walk(sh.shapes):
                if getattr(s, "has_text_frame", False):
                    names.update(TOKEN.findall(s.text_frame.text))
            names &= present
            if names and (all(not value(n) for n in names) or
                          any(rule(n)[1] == "remove_row" and not value(n) for n in names)):
                _remove(sh._element)
            else:
                prune(sh.shapes)

    for sl in slides:
        prune(sl.shapes)

    # 4. fill; hide or clear empty lines
    for sl in slides:
        for sh in _text_shapes(sl):
            tf = sh.text_frame
            for p in list(tf.paragraphs):
                ptext = "".join(r.text for r in p.runs)
                names = [n for n in TOKEN.findall(ptext) if n in present]
                if not names:
                    continue
                if all(not value(n) for n in names):
                    pols = {rule(n)[1] for n in names}
                    if pols == {"hide_line"} and len(tf.paragraphs) > 1:
                        _remove(p._p)
                        continue
                    if pols == {"clear"}:
                        for r in p.runs:
                            r.text = ""
                        continue
                for r in p.runs:
                    r.text = TOKEN.sub(
                        lambda m: (value(m.group()) or FIRM_TODO) if m.group() in present else m.group(),
                        r.text)

    # 5. safety net: no template variable may survive
    for sl in slides:
        for sh in _text_shapes(sl):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if any(m.group() in present for m in TOKEN.finditer(r.text)):
                        r.text = TOKEN.sub(lambda m: FIRM_TODO if m.group() in present else m.group(), r.text)
                        warnings.append("a variable survived filling and was replaced")

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), warnings


def template_variables(template=TEMPLATE):
    """The variable names the template holds (for checks and tests)."""
    prs = _load_template(template)
    names = set()
    for si, sl in enumerate(prs.slides, 1):
        for sh in _text_shapes(sl):
            for p in sh.text_frame.paragraphs:
                _join_split_tokens(p, f"slide {si}")
                for r in p.runs:
                    names.update(TOKEN.findall(r.text))
    return names


if __name__ == "__main__":
    import json
    import sys
    payload = json.load(open(sys.argv[1]))
    if "values" not in payload:
        payload = {"values": payload}
    pptx, warns = build_deck(payload["values"], payload.get("rules"))
    dest = sys.argv[2] if len(sys.argv) > 2 else "out.pptx"
    open(dest, "wb").write(pptx)
    print("wrote", dest, "| warnings:", warns or "none")
