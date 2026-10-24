"""S3 deck layout (template "1.5"): template 1.4 filled as usual, then restructured.

- Section dividers are left out.
- The executive summary stays the first slide after the cover and the agenda (José, 24 Oct).
- The template's own timeline ("Plan by workstream", 1-2 pages) stays in the story, right after the approach
  overview (José, 24 Oct: slide 7 keeps the 1.4 timeline's structure; he will explain how to use it for long
  plans). The custom one-page roadmaps tried on 23-24 Oct are gone.
- The appendix comes after the closing slide, for an executive reader (José, 24 Oct): "Deliverables by
  workstream" (the steering-meetings slide refers readers to it), then the assumptions the draft rests on
  (payload s3.assumptions), which explain our rationale to readers who want it.
- The timeline uses at most 24 boxes of one unit (week, month, bi-month, quarter, four-month period, semester;
  workplan.time_unit). A plan longer than 24 semesters gets no plan slides at all (José, 24 Oct).
- New slide: a "[Firm to complete] Fees" placeholder before Next Steps (no team slide and no price: the firm
  completes both).
- The agenda follows the new order; the situation slide's box reads "Hypothesis to test".

Called by varfill.build_deck after filling, before saving. Payloads without template "1.5" never reach it.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import workplan as wp

NAVY = RGBColor(0x14, 0x2B, 0x4C)
TEAL = RGBColor(0x2E, 0x86, 0x8C)
GREY = RGBColor(0xE6, 0xE9, 0xEE)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

X0, Y0 = 554736, 1700000      # content area under the subtitle (Default layout of template 1.4)
YMAX = 6250000                # above the footer
FEES_TITLE = "[Firm to complete] Fees"
NUM_WORDS = ["no", "one", "two", "three", "four", "five", "six", "seven", "eight"]


def _title(s):
    for sh in s.placeholders:
        if sh.placeholder_format.idx == 0:
            return sh.text_frame.text.strip()
    return ""


def _texts(s):
    out = []
    for sh in s.shapes:
        out.extend(_shape_texts(sh))
    return " | ".join(out)


def _shape_texts(sh):
    if sh.shape_type == 6:  # group
        return [t for g in sh.shapes for t in _shape_texts(g)]
    return [sh.text_frame.text] if getattr(sh, "has_text_frame", False) and sh.has_text_frame else []


def _weeks(d):
    d = d or {}
    lo, hi = d.get("min"), d.get("max")
    return f"{hi}" if lo == hi else f"{lo}–{hi}"


class _Builder:
    def __init__(self, prs):
        self.prs = prs
        self.W = prs.slide_width
        self.CW = self.W - 2 * X0
        self.layout = [lo for lo in prs.slide_layouts if lo.name == "Default"][0]

    def slide(self, title, subtitle, label=""):
        s = self.prs.slides.add_slide(self.layout)
        for ph in s.placeholders:
            idx = ph.placeholder_format.idx
            ph.text_frame.text = {0: title, 1: subtitle, 17: label}.get(idx, "")
        return s

    def box(self, s, x, y, w, h, text, size=12, bold=False, color=INK, fill=None, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, shape=MSO_SHAPE.RECTANGLE, tight=False):
        sh = s.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(max(int(w), 1)), Emu(max(int(h), 1)))
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(72000)
        tf.margin_top = tf.margin_bottom = Emu(0 if tight else 36000)
        for i, ln in enumerate(text if isinstance(text, list) else [text]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        return sh

    def line(self, s, x1, y1, x2, y2, color, width=1.0):
        ln = s.shapes.add_connector(1, Emu(int(x1)), Emu(int(y1)), Emu(int(x2)), Emu(int(y2)))
        ln.line.color.rgb = color
        ln.line.width = Pt(width)
        return ln

    # ---------- assumptions ----------
    def assumptions(self, items, label):
        n = len(items)
        s = self.slide(f"This draft rests on {NUM_WORDS[n] if n < len(NUM_WORDS) else n} assumption{'s' if n != 1 else ''} "
                       f"to confirm before the work starts", "Assumptions we made", label)
        y = Y0
        self.box(s, X0 + 900000, y, 5600000, 300000, "What we assumed", size=11, bold=True, color=MUTED)
        self.box(s, X0 + 6600000, y, self.CW - 6600000, 300000, "If it is wrong", size=11, bold=True, color=MUTED)
        y += 340000
        row = min(660000, (YMAX - y) // max(1, n))
        for a in items:
            self.box(s, X0, y, 800000, row - 100000, str(a.get("id", "")), size=18, bold=True, color=TEAL,
                     anchor=MSO_ANCHOR.MIDDLE)
            self.box(s, X0 + 900000, y, 5600000, row - 100000, str(a.get("text", "")), size=14,
                     anchor=MSO_ANCHOR.MIDDLE)
            self.box(s, X0 + 6600000, y, self.CW - 6600000, row - 100000, str(a.get("if_wrong", "")), size=12,
                     color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
            self.line(s, X0, y + row - 60000, X0 + self.CW, y + row - 60000, GREY)
            y += row
        return s

    # ---------- fees placeholder ----------
    def fees(self, plan, label):
        s = self.slide(FEES_TITLE, "Fees and commercial terms", label)
        self.box(s, X0, Y0, self.CW, 1300000,
                 ["[Firm to complete] Add the fees and commercial terms.", "This draft contains no price."],
                 size=20, bold=True, color=NAVY, fill=GREY, anchor=MSO_ANCHOR.MIDDLE)
        phases = (plan or {}).get("phases") or []
        if phases:
            span = (plan or {}).get("total_duration_weeks") or {}
            parts = ", ".join(f"{p['name']} {_weeks(p.get('duration_weeks'))} weeks" for p in phases)
            self.box(s, X0, Y0 + 1450000, self.CW, 600000,
                     f"For pricing, the plan runs {_weeks(span)} weeks: {parts}.", size=12, color=MUTED)
        return s


def apply(prs, values, plan, s3):
    """Restructure a filled template-1.4 deck into the S3 layout. Returns warnings."""
    warnings = []
    values = values or {}
    s3 = s3 or {}
    plan = plan if isinstance(plan, dict) else {}
    sec = [str(values.get(f"section_{i:02d}") or "").strip() for i in range(1, 7)]
    sec_names = {x for x in sec if x}
    b = _Builder(prs)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    slides = list(prs.slides)
    kinds = {}
    for el, s in zip(ids, slides):
        t, tx, lay = _title(s), _texts(s), s.slide_layout.name
        if lay == "Top left title" and t in sec_names:
            kinds[el] = "divider"
        elif lay == "Top left title" and "section_" not in tx and any(x in tx for x in sec_names) and len(sec_names) > 1 \
                and sum(1 for x in sec_names if x in tx) >= 2:
            kinds[el] = "agenda"
        elif "Objectives" in tx and "Outcomes" in tx:
            kinds[el] = "summary"
        elif "Where things stand" in tx:
            kinds[el] = "situation"
        elif "Deliverables by workstream" in tx:
            kinds[el] = "deliverables"
        elif "Plan by workstream" in tx:
            kinds[el] = "workplan"
        elif "Approach overview" in tx:
            kinds[el] = "overview"
        elif "Steering meetings" in tx:
            kinds[el] = "steering"
        elif lay == "Default" and t and t == (sec[3] or "Next Steps"):
            kinds[el] = "next"
        else:
            kinds[el] = "keep"
    label_approach = sec[2] or "Approach"
    label_next = sec[3] or "Next Steps"

    # new slides (appended; placed below)
    has_plan = bool(plan.get("phases"))
    # José, 24 Oct: a plan longer than 24 semesters sells nothing for the short term. Every other slide is
    # produced; the Phase Architect's slides are not, and the reply email says why (Build render payload).
    weeks = int((plan.get("total_duration_weeks") or {}).get("max") or 0)
    too_long = has_plan and weeks > wp.MAX_WEEKS
    if too_long:
        warnings.append(f"S3: the plan runs {weeks} weeks, more than 24 semesters; the plan slides are left out")
    items = [a for a in (s3.get("assumptions") or []) if str(a.get("text", "")).strip()]
    assum = b.assumptions(items[:5], "Appendix") if items else None
    fees = b.fees(plan, label_next)
    if not items:
        warnings.append("S3: no assumptions in the payload; the assumptions slide is left out")
    new = {el: k for el, k in zip(list(lst)[len(ids):],
                                  [k for k, v in (("assumptions", assum), ("fees", fees)) if v is not None])}
    by_kind = {k: el for el, k in new.items()}

    order, appendix = [], []
    next_seen = False
    for e in ids:
        k = kinds[e]
        if k == "divider":
            continue
        if too_long and k in ("overview", "workplan", "deliverables", "steering"):
            continue
        if k == "deliverables" and has_plan:
            # appendix (José, 24 Oct: executive level); the steering-meetings slide sends readers to it
            appendix.append(e)
            continue
        if k == "next":
            order += [by_kind[x] for x in ("fees",) if x in by_kind]
            next_seen = True
        order.append(e)
    if not next_seen:                                 # no Next Steps slide: new slides go before the closing slide
        tail = order.pop() if order else None
        order += [by_kind[x] for x in ("fees",) if x in by_kind]
        if tail is not None:
            order.append(tail)
    if "assumptions" in by_kind:
        appendix.append(by_kind["assumptions"])       # last: after the deliverables (José, 24 Oct)
    if appendix:
        order = order + appendix                      # after the closing slide (José, 24 Oct)
        for e in appendix:                            # the tracker (top right) reads "Appendix" on every one
            sl = prs.slides.get(int(e.get("id")))
            for ph in (sl.placeholders if sl is not None else []):
                if ph.placeholder_format.idx == 17:
                    ph.text_frame.text = "Appendix"

    for e in list(lst):
        lst.remove(e)
        if e not in order:
            prs.part.drop_rel(e.get(qn("r:id")))
    for e in order:
        lst.append(e)

    # agenda in the new order; situation box label
    for s in prs.slides:
        for sh in s.shapes:
            for t in _iter_text_shapes(sh):
                for p in t.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if txt.strip() == "What we believe" and p.runs:
                        p.runs[0].text = "Hypothesis to test"
                        for r in p.runs[1:]:
                            r.text = ""
                    elif txt.startswith("Hypothesis to test:") and p.runs:
                        # the box is now headed "Hypothesis to test": the line drops the repeated label
                        rest = txt[len("Hypothesis to test:"):].strip()
                        p.runs[0].text = rest[:1].upper() + rest[1:]
                        for r in p.runs[1:]:
                            r.text = ""
    _reorder_agenda(prs, ids, kinds, sec, False)     # assumptions sit in the appendix, not under Next Steps
    return warnings


def _iter_text_shapes(sh):
    if sh.shape_type == 6:
        for g in sh.shapes:
            yield from _iter_text_shapes(g)
    elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        yield sh


def _reorder_agenda(prs, ids, kinds, sec, has_assumptions):
    names = [x for x in sec if x]
    if len(names) < 2:
        return
    new = list(names)                                 # Executive Summary stays first (José, 24 Oct)
    if has_assumptions and len(names) >= 4 and names[3]:
        new[3] = f"Assumptions and {names[3]}" if not names[3].lower().startswith("assumptions") else names[3]
    agenda = None
    for e, s in zip(list(prs.slides._sldIdLst), prs.slides):
        if kinds.get(e) == "agenda":
            agenda = s
    if agenda is None:
        return
    seen = 0
    for sh in agenda.shapes:
        for t in _iter_text_shapes(sh):
            for p in t.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs).strip()
                if txt in names and seen < len(new) and p.runs:
                    p.runs[0].text = new[seen]
                    for r in p.runs[1:]:
                        r.text = ""
                    seen += 1
