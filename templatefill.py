"""
templatefill.py — JML AI Analyst deck renderer (v2, template-based).

Opens the bundled 9-slide JML template and fills the named shapes from the
deck JSON payload, preserving each shape's own formatting (fonts, colours,
positions). Static slides (agenda, dividers, closing) are left untouched.

build_deck(data: dict) -> bytes   # returns .pptx bytes

Fill plan (confirmed with Jose, AI_Analyst_Template_Fill_Plan.md):
  Slide 1 (idx0)  cover      : keep "Consulting Proposal", replace subtitle run
  Slide 2 (idx1)  agenda     : STATIC
  Slide 3 (idx2)  divider    : STATIC
  Slide 4 (idx3)  context    : Title + Objectives (from brief);
                               Solution + Outcomes = [Firm to complete]
  Slide 5 (idx4)  divider    : STATIC
  Slide 6 (idx5)  market     : title + 4 driver boxes (findings) +
                               4 rationale boxes (implications) + source line
  Slide 7 (idx6)  divider    : STATIC
  Slide 8 (idx7)  next steps : Title + up to 6 step boxes
  Slide 9 (idx8)  closing    : STATIC
"""

import os
from io import BytesIO
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "template.pptx")
FIRM_TODO = "[Firm to complete]"


# ----------------------------------------------------------------- helpers
def _shorten(s, n):
    """Trim to n chars at a word boundary, add an ellipsis if trimmed."""
    s = " ".join(str(s or "").split()).strip()
    if len(s) <= n:
        return s
    cut = s[: n - 1]
    sp = cut.rfind(" ")
    if sp > n * 0.6:          # only break on a space if it isn't too early
        cut = cut[:sp]
    return cut.rstrip(" ,.;:") + "…"


def _find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def _set_para(para, text):
    """Put `text` in the paragraph's first run; blank any other runs.
    Preserves the first run's font (size/bold/colour)."""
    runs = para.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    elif text:
        para.add_run().text = text


def _fill_one(shape, text):
    """Single-line/flow box: write text into para 0, clear the rest."""
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    _set_para(tf.paragraphs[0], text)
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


def _fill_paras(shape, items):
    """Multi-paragraph box: map items onto existing paragraphs in order,
    respecting the box's designed paragraph count (extra items are dropped,
    extra paragraphs are blanked). Keeps each paragraph's own formatting."""
    if shape is None or not shape.has_text_frame:
        return
    paras = shape.text_frame.paragraphs
    for i, p in enumerate(paras):
        _set_para(p, items[i] if i < len(items) else "")


def _remove(shape):
    """Delete a shape from its slide/group."""
    if shape is not None:
        shape._element.getparent().remove(shape._element)


def _step_rect(group):
    """Return the text rectangle inside a Next-step group (the non-numeric one)."""
    for sub in group.shapes:
        if sub.has_text_frame:
            t = sub.text_frame.text.strip()
            if t and not t.isdigit():
                return sub
    return None


# ----------------------------------------------------------------- builder
def build_deck(data):
    data = data or {}
    cover = data.get("cover") or {}
    context = data.get("context") or {}
    market = data.get("market_research") or {}
    nexts = data.get("next_steps") or {}
    rows = market.get("rows") or []
    sources = market.get("sources") or []

    client = (cover.get("title") or "").replace("Draft proposal for ", "").strip() or "the client"
    subject = cover.get("subtitle") or ""

    prs = Presentation(TEMPLATE)
    S = list(prs.slides)

    # --- Slide 1: cover ---------------------------------------------------
    sh = _find(S[0], "Google Shape;187;p1")
    if sh and sh.has_text_frame:
        runs = sh.text_frame.paragraphs[0].runs
        if len(runs) >= 2:
            runs[1].text = _shorten("A proposal for " + client, 90)
            for r in runs[2:]:
                r.text = ""

    # --- Slide 4: context (Objectives filled; Solution/Outcomes = TODO) ---
    _fill_one(_find(S[3], "Title 1"),
              _shorten(subject or (client + " — situation and objectives"), 90))

    obj_items = [_shorten(b, 130) for b in (context.get("bullets") or [])][:4]
    _fill_paras(_find(S[3], "Google Shape;881;p247"),
                obj_items or [FIRM_TODO])

    _fill_paras(_find(S[3], "Google Shape;882;p247"),
                [FIRM_TODO + " The proposed approach and phases are added once the plan is confirmed."])
    _fill_paras(_find(S[3], "Google Shape;883;p247"),
                [FIRM_TODO + " Expected results and deliverables are added with the plan."])

    # --- Slide 6: market research ----------------------------------------
    _fill_one(_find(S[5], "Rectangle 3"),
              _shorten(market.get("title") or "What peers and the market are doing", 90))

    n = len(rows)
    line_above = {1: "Line 17", 2: "Line 18", 3: "Line 19"}  # dashed row separators
    for i in range(4):
        dbox = _find(S[5], f"AutoShape {4 + i}")
        rbox = _find(S[5], f"Rectangle {11 + i}")
        if i < n:
            r = rows[i]
            _fill_one(dbox, _shorten(r.get("finding") or r.get("actor") or "", 72))
            bullets = []
            if r.get("implication"):
                bullets.append(_shorten(r["implication"], 150))
            if r.get("actor"):
                bullets.append("— " + _shorten(r["actor"], 60))
            _fill_paras(rbox, bullets or [""])
        elif i == 0 and n == 0:
            # no findings at all: drop the empty driver arrow, keep an honest note
            _remove(dbox)
            _fill_paras(rbox, ["No market findings were captured for this brief."])
        else:
            # unused slot: remove the box pair and the separator line above it
            _remove(dbox)
            _remove(rbox)
            _remove(_find(S[5], line_above.get(i)))

    # "Upsides not included in business case" row has no market mapping -> placeholder.
    # Slide it (and its separator) up to follow the last real row so there is no gap.
    _fill_paras(_find(S[5], "Rectangle 15"), [FIRM_TODO])
    vis = max(n, 1)                       # visible rows above the upsides row
    up_top = Inches(2.24 + vis * 0.92)
    for nm, off in (("AutoShape 8", 0.0), ("Rectangle 15", 0.0), ("Line 20", -0.04)):
        sh2 = _find(S[5], nm)
        if sh2 is not None:
            sh2.top = up_top + Inches(off)

    src = _find(S[5], "Google Shape;507;p14")
    if src:
        if sources:
            line = "Sources: " + "   ".join(f"[{i+1}] {u}" for i, u in enumerate(sources))
            _fill_one(src, _shorten(line, 180))
        else:
            _fill_one(src, "")

    # --- Slide 8: next steps (6 step boxes, top -> bottom) ----------------
    _fill_one(_find(S[7], "Title 1"), nexts.get("title") or "Next steps")
    steps = [_shorten(b, 180) for b in (nexts.get("bullets") or [])]
    groups = sorted(
        [sh for sh in S[7].shapes if sh.shape_type == MSO_SHAPE_TYPE.GROUP],
        key=lambda g: g.top or 0,
    )
    for i, g in enumerate(groups):
        rect = _step_rect(g)
        if rect is None:
            continue
        _fill_one(rect, steps[i] if i < len(steps) else FIRM_TODO)
    _fill_one(_find(S[7], "Google Shape;507;p14"), "")   # no sources on next steps

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    import json, sys
    payload = json.load(open(sys.argv[1])) if len(sys.argv) > 1 else {}
    out = build_deck(payload)
    dest = sys.argv[2] if len(sys.argv) > 2 else "out.pptx"
    open(dest, "wb").write(out)
    print(f"wrote {dest} ({len(out)} bytes)")
