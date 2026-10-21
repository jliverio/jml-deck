"""Render every plan written by phase_architect/tests/plan_size_test.js (validator-compiled, real repairs and
status weeks) and run the layout checks of tests/workplan_stress.py on each slide.
Usage: node phase_architect/tests/plan_size_test.js && python3 tests/workplan_render_size.py [outdir]"""
import glob, io, json, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
import workplan
from workplan_stress import check_slide

SRC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "phase_architect", "tests", "out_plan_size")
outdir = sys.argv[1] if len(sys.argv) > 1 else None
bad = 0
files = sorted(glob.glob(os.path.join(SRC, "*.json")))
for f in files:
    plan = json.load(open(f))
    name = os.path.basename(f)[:-5]
    data = workplan.build_workplan_deck(plan, client="Testco")
    prs = Presentation(io.BytesIO(data))
    probs = []
    for i, sl in enumerate(prs.slides, 1):
        probs += check_slide(sl, prs.slide_width, prs.slide_height, i)
    if outdir:
        os.makedirs(outdir, exist_ok=True)
        open(os.path.join(outdir, name + ".pptx"), "wb").write(data)
    bad += bool(probs)
    print(("FAIL " if probs else "PASS ") + f"{name}: {len(prs.slides)} slides" + ("" if not probs else f"; " + " | ".join(probs[:3])))
print(f"\n{bad} failing of {len(files)}")
