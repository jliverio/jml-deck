"""Stress the work-plan slides across durations and stream counts.

    python3 tests/workplan_stress.py [outdir]

Builds synthetic compiled plans (same shape as validate_plan_v2.js output)
for a grid of durations x workstream counts, renders them with workplan.py,
and checks the layout: every shape inside the slide, no text box whose text
cannot fit, no overlapping labels. Prints one line per case; exits 1 on any
failure. With outdir, also writes each deck there for visual review.
"""
import io, json, math, os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
from pptx import Presentation
from pptx.util import Emu
import workplan

WS_NAMES = ["Customer needs and demand", "Distributor channel risks", "Prior web shop lessons",
            "Direct channel options", "Direct channel model", "Safeguards and KPIs",
            "Pilot plan and final deck", "Pricing and assortment", "Operations and logistics readiness",
            "Change management and training", "Data and IT platform requirements", "Business case"]
DEL_NAMES = ["Customer needs and demand evidence", "Distributor volume and relationship risk map",
             "Lessons from the 2023 web shop", "Direct channel options with fit assessment",
             "Direct channel blueprint with distributor roles", "Agreed KPIs and distributor safeguards",
             "Pilot plan", "Final presentation deck", "Pricing and assortment recommendation",
             "Logistics readiness assessment", "Training plan for the sales team", "Business case and roadmap"]
PHASE_NAMES = ["Research", "Design", "Pilot", "Scale-up"]


def make_plan(total, n_ws, n_phases=None, dels_per_ws=1, cadence=2):
    n_phases = n_phases or (1 if total <= 4 else (2 if n_ws <= 8 or total < 16 else 3))
    n_phases = min(n_phases, n_ws, total)
    # split weeks across phases (first phase a bit shorter)
    base = [total // n_phases] * n_phases
    for i in range(total - sum(base)):
        base[-1 - i] += 1
    phases, start = [], 1
    for i, d in enumerate(base):
        phases.append({"name": PHASE_NAMES[i], "order": i + 1, "start_week": start, "end_week": start + d - 1,
                       "duration_weeks": {"min": d, "max": d}, "workstream_ids": [],
                       "gate": {"label": f"{PHASE_NAMES[i]} sign-off", "decision": "Go on?"}})
        start += d
    ws, dels = [], []
    per = [n_ws // n_phases + (1 if i < n_ws % n_phases else 0) for i in range(n_phases)]
    k = 0
    for p, cnt in zip(phases, per):
        L = p["end_week"] - p["start_week"] + 1
        for j in range(cnt):
            s = p["start_week"] + (j * L) // (cnt + 1)
            e = min(p["end_week"], s + max(1, (2 * L) // 3) - 1)
            wid = f"W{k + 1}"
            ids = []
            for m in range(dels_per_ws):
                did = f"D{len(dels) + 1}"
                due = max(s, e - (dels_per_ws - 1 - m))
                dels.append({"id": did, "name": DEL_NAMES[len(dels) % len(DEL_NAMES)], "kind": "analysis",
                             "workstream_id": wid, "phase_order": p["order"], "due_week": due, "presented_at": None})
                ids.append(did)
            ws.append({"id": wid, "phase_order": p["order"], "name": WS_NAMES[k % len(WS_NAMES)],
                       "start_week": s, "end_week": e, "deliverable_ids": ids})
            p["workstream_ids"].append(wid)
            k += 1
    events = [{"id": "G1", "type": "kickoff", "phase_order": 1, "week": 1, "label": "Kick-off",
               "purpose": "Align on scope and plan", "attendees": ["CCO", "Head of Digital"], "presents": [],
               "decision": None, "agenda": ["Present the plan", "Confirm stakeholders"]}]
    for p in phases:
        L = p["end_week"] - p["start_week"] + 1
        if L >= 4:
            wk = p["start_week"] + L // 2 - 1
            events.append({"id": f"G{len(events)+1}", "type": "checkin", "phase_order": p["order"], "week": wk,
                           "label": f"Check-in: {p['name'].lower()} progress", "purpose": "Review early findings",
                           "attendees": ["CCO"], "presents": [], "decision": None,
                           "agenda": ["Present progress", "Discuss open questions"]})
        events.append({"id": f"G{len(events)+1}", "type": "gate", "phase_order": p["order"], "week": p["end_week"],
                       "label": p["gate"]["label"], "purpose": f"Sign off the {p['name'].lower()} phase",
                       "attendees": ["CCO", "Head of Digital"], "presents": [], "decision": "Does the work go forward?",
                       "agenda": ["Present results", "Decide next phase"]})
    by_week = {}
    for d in dels:  # present each deliverable at the first event on or after its due week
        ev = next((e for e in sorted(events, key=lambda e: e["week"]) if e["week"] >= d["due_week"] and e["type"] != "kickoff"), events[-1])
        d["presented_at"] = ev["label"]; ev["presents"].append(d["id"])
    taken = {e["week"] for e in events}
    status = [w for w in range(1 + cadence, total + 1, cadence) if w not in taken]
    return {"total_duration_weeks": {"min": total, "max": total}, "phases": phases, "workstreams": ws,
            "deliverables": dels, "deadline": None,
            "governance": {"status_cadence_weeks": cadence, "status_weeks": status, "events": events}}


# ----------------------------------------------------------- layout checks
def _avg_char_w(size_pt, bold):
    return size_pt * (0.56 if bold else 0.52) * 12700  # EMU per char, Arial approx


def text_need(shape):
    """Estimated EMU height the shape's text needs at its width."""
    tf = shape.text_frame
    w = shape.width - (tf.margin_left or 0) - (tf.margin_right or 0)
    if w <= 0:
        return 10**9
    h = 0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not p.runs:
            continue
        f = p.runs[0].font
        size = f.size.pt if f.size else 12
        bold = bool(f.bold)
        indent = 171450 if p._p.pPr is not None and p._p.pPr.get("marL") else 0
        per_line = max(1, int((w - indent) / _avg_char_w(size, bold)))
        words, lines, cur = txt.split(), 1, 0
        for wd in words:
            add = len(wd) + (1 if cur else 0)
            if cur + add > per_line and cur:
                lines += 1; cur = len(wd)
            else:
                cur += add
        h += lines * size * 1.2 * 12700 + (p.space_after.pt * 12700 if p.space_after else 0)
    return h + (tf.margin_top or 0) + (tf.margin_bottom or 0)


def text_span(shape):
    """Horizontal extent the text actually covers (box width caps it)."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    f = p.runs[0].font if p.runs else None
    size = f.size.pt if f is not None and f.size else 12
    longest = max((len("".join(r.text for r in q.runs)) for q in tf.paragraphs), default=0)
    w = min(shape.width, int(longest * _avg_char_w(size, bool(f and f.bold))))
    al = p.alignment
    if al is not None and int(al) == 2:      # centre
        x0 = shape.left + (shape.width - w) // 2
    elif al is not None and int(al) == 3:    # right
        x0 = shape.left + shape.width - w
    else:
        x0 = shape.left
    return x0, x0 + w


def check_slide(sl, sw, sh, idx):
    probs = []
    labels = []
    for s in sl.shapes:
        if s.left is None:
            continue
        if s.left < -Emu(9144) or s.top < -Emu(9144) or s.left + s.width > sw + Emu(9144) or s.top + s.height > sh + Emu(9144):
            probs.append(f"s{idx}: '{s.name}' outside slide")
        if getattr(s, "has_text_frame", False) and s.text_frame.text.strip() and not s.is_placeholder:
            need = text_need(s)
            if need > s.height * 1.15 + 12700 * 2:
                probs.append(f"s{idx}: text overflows ({int(need/12700)}pt in {int(s.height/12700)}pt): {s.text_frame.text[:40]!r}")
            labels.append(s)
    # overlapping free-standing text boxes (not text inside filled shapes)
    tbs = [s for s in labels if s.shape_type == 17]
    for i in range(len(tbs)):
        for j in range(i + 1, len(tbs)):
            a, b = tbs[i], tbs[j]
            ax0, ax1 = text_span(a); bx0, bx1 = text_span(b)
            ox = min(ax1, bx1) - max(ax0, bx0)
            oy = min(a.top + text_need(a), b.top + text_need(b)) - max(a.top, b.top)
            if ox > 12700 * 3 and oy > 12700 * 3:
                probs.append(f"s{idx}: labels overlap: {a.text_frame.text[:25]!r} / {b.text_frame.text[:25]!r}")
    return probs


def run_case(total, n_ws, dels_per_ws=1, outdir=None):
    plan = make_plan(total, n_ws, dels_per_ws=dels_per_ws)
    data = workplan.build_workplan_deck(plan, client="The Navigator Company")
    prs = Presentation(io.BytesIO(data))
    probs = []
    for i, sl in enumerate(prs.slides, 1):
        probs += check_slide(sl, prs.slide_width, prs.slide_height, i)
    name = f"wp_{total:02d}w_{n_ws:02d}ws_{dels_per_ws}d"
    if outdir:
        open(os.path.join(outdir, name + ".pptx"), "wb").write(data)
    return name, len(prs.slides), probs


CASES = [(t, n, d) for t in (4, 6, 12, 24, 40, 52) for n in (3, 7, 12) for d in (1, 2)]

if __name__ == "__main__":
    outdir = sys.argv[1] if len(sys.argv) > 1 else None
    bad = 0
    for t, n, d in CASES:
        if n > t * 2:
            continue
        name, ns, probs = run_case(t, n, d, outdir)
        bad += bool(probs)
        print(("FAIL " if probs else "PASS ") + f"{name}: {ns} slides" + ("" if not probs else f", {len(probs)} problems; e.g. " + " | ".join(probs[:3])))
    print(f"\n{bad} failing cases")
    sys.exit(1 if bad else 0)
