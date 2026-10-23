"""Usage: node phase_architect/tests/plan_size_test.js && python3 tests/s3_roadmap_stress.py
Draw the S3 roadmap for every plan-size fixture (1-5 phases, 2-52 weeks, up to 12 workstreams) with and
without a deadline; flag shapes outside the content area and text that likely overflows its box."""
import json, glob, math, sys, copy, os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import varfill, s3layout
from pptx.util import Pt
bad = []; n = 0; ws_hidden = 0
for f in sorted(glob.glob(os.path.join(ROOT, 'phase_architect', 'tests', 'out_plan_size', '*.json'))):
    plan = json.load(open(f))
    for dl in (None, (plan['total_duration_weeks'] or {}).get('max')):
        prs = varfill._load_template(varfill.TEMPLATES['1.4'])
        b = s3layout._Builder(prs)
        p = copy.deepcopy(plan)
        for ph in p['phases']:
            ph['gate'] = {'label': 'Go/no-go on the pilot'[:22], 'decision': 'x'}
        s, show = b.roadmap(p, dl, 'Approach'); n += 1; ws_hidden += (not show)
        for sh in s.shapes:
            if sh.is_placeholder: continue
            if sh.top + sh.height > 6400000 or sh.left + sh.width > b.W + 1000 or sh.left < 0:
                bad.append((f.split('/')[-1], dl, 'out of area', sh.text_frame.text[:30] if sh.has_text_frame else ''))
            if sh.has_text_frame and sh.text_frame.text.strip():
                size = sh.text_frame.paragraphs[0].runs[0].font.size.pt
                cw = size * 0.5 * 12700; lh = size * 1.2 * 12700
                usable = max(1, sh.width - 144000)
                words = sh.text_frame.text.split(); lines = 1; cur = 0
                for w in words:
                    wl = (len(w) + 1) * cw
                    if cur + wl > usable and cur > 0: lines += 1; cur = wl
                    else: cur += wl
                if lines * lh + (sh.text_frame.margin_top or 0) + (sh.text_frame.margin_bottom or 0) > sh.height + 20000:
                    bad.append((f.split('/')[-1], dl, f'text {lines} lines in {sh.height}', sh.text_frame.text[:40]))
print('roadmaps', n, 'workstreams moved to appendix', ws_hidden, 'problems', len(bad))
from collections import Counter
print(Counter(b[2].split(' ')[0] + ' ' + b[2].split(' ')[1] for b in bad).most_common(8))
for x in bad[:25]: print(' ', x)

sys.exit(1 if bad else 0)
