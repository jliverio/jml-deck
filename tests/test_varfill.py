"""Offline tests for varfill.py against template 1.2 and the seed table.

Run:  python3 tests/test_varfill.py
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from pptx import Presentation  # noqa: E402
import varfill  # noqa: E402
import workplan as wp  # noqa: E402

RULES = json.load(open(os.path.join(ROOT, "variables", "template_variables.json")))
METALCO = json.load(open(os.path.join(HERE, "metalco.json")))
VAR_RE = re.compile(r"(?<![A-Za-z0-9_])(cover|section|finding|sources|next_step|summary|objective|solution|outcome)_[a-z0-9_]+")


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def texts(pptx_bytes):
    prs = Presentation(io.BytesIO(pptx_bytes))
    out = []
    for i, sl in enumerate(prs.slides, 1):
        for sh in walk(sl.shapes):
            if getattr(sh, "has_text_frame", False):
                out.append((i, sh.name, sh.text_frame.text))
    return prs, out


def groups(prs, slide_no):
    return sum(1 for sh in list(prs.slides)[slide_no - 1].shapes if sh.shape_type == 6)


results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))


# 1. template and table agree
tv = varfill.template_variables()
table_names = {r["variable"] for r in RULES}
check("every template variable has a table row", tv <= table_names, sorted(tv - table_names))

# 2. realistic deck with the real rules
pptx, warns = varfill.build_deck(METALCO, RULES)
prs, tx = texts(pptx)
left = [t for t in tx if VAR_RE.search(t[2])]
check("no variable names left in the deck", not left, left[:3])
check("agenda shows 3 topics", groups(prs, 2) == 3, groups(prs, 2))
check("market slide shows 3 finding rows", groups(prs, 6) == 3, groups(prs, 6))
check("next steps shows 3 steps", groups(prs, 8) == 3, groups(prs, 8))
cover = [t for t in tx if t[0] == 1 and "Consulting Proposal" in t[2]]
check("cover title and subtitle filled separately",
      cover and "A proposal for MetalCo" in cover[0][2] and "sub" not in cover[0][2].replace("subtitle", ""),
      cover)
f2 = [t for t in tx if t[0] == 6 and "Eurostat reported" in t[2]]
check("empty result line hidden (finding 2 has 2 lines)",
      f2 and len(f2[0][2].split("\n")) == 2, f2[0][2] if f2 else None)
src = [t for t in tx if t[0] == 6 and t[2].startswith("Sources:")]
check("sources footer filled", src and "[1] Hierros" in src[0][2], src)
check("no cut warnings on realistic input", not [w for w in warns if " cut to " in w], warns)
# the table also holds template 1.4's 10 rows (situation slide, footer, work-plan titles), absent from 1.2
check("table rows beyond 1.2 are only 1.4's", not any("table rows have no variable" in w for w in warns)
      or any(w.startswith("10 table rows have no variable") for w in warns), warns)

# 3. empty action removes the whole finding row (remove_row)
v = dict(METALCO)
v["finding_03_action_taken"] = ""
prs3, _ = texts(varfill.build_deck(v, RULES)[0])
check("empty action removes the finding row", groups(prs3, 6) == 2, groups(prs3, 6))

# 4. empty sources clears the whole footer line, label included (clear)
v = dict(METALCO)
v["sources_market_01"] = ""
_, tx4 = texts(varfill.build_deck(v, RULES)[0])
check("empty sources clears 'Sources:' label", not any(t[2].startswith("Sources") for t in tx4))

# 5. no findings at all: market slide has no rows, deck still renders
v = {k: x for k, x in METALCO.items() if not k.startswith("finding_")}
prs5, tx5 = texts(varfill.build_deck(v, RULES)[0])
check("no findings -> no finding rows", groups(prs5, 6) == 0, groups(prs5, 6))

# 6. too many agenda topics stops the render
v = dict(METALCO)
v["section_08"] = "Appendix"
try:
    varfill.build_deck(v, RULES)
    check("8 topics raise TemplateCapacityError", False)
except varfill.TemplateCapacityError as e:
    check("8 topics raise TemplateCapacityError", "8 agenda topics" in str(e), str(e))

# 7. a template variable missing from the table stops the render
try:
    varfill.build_deck(METALCO, [r for r in RULES if r["variable"] != "finding_02_context"])
    check("missing table row raises TemplateRulesError", False)
except varfill.TemplateRulesError as e:
    check("missing table row raises TemplateRulesError", "finding_02_context" in str(e), str(e))

# 8. limits come from the table, not the code
rules = [dict(r) for r in RULES]
for r in rules:
    if r["variable"] == "next_step_01":
        r["max_chars"] = 20
pptx8, w8 = varfill.build_deck(METALCO, rules)
_, tx8 = texts(pptx8)
ns = [t[2] for t in tx8 if t[0] == 8 and t[2].startswith("Confirm")]
check("table max_chars is enforced", ns and len(ns[0]) <= 20 and ns[0].endswith("…"), ns)
check("cut produces a warning", any("next_step_01" in w for w in w8), w8)

# 9. a long source list is cut at a whole entry
v = dict(METALCO)
v["sources_market_01"] = "; ".join(f"[{i}] Publisher number {i}, a long report title, 2025" for i in range(1, 9))
_, tx9 = texts(varfill.build_deck(v, RULES)[0])
s9 = [t[2] for t in tx9 if t[0] == 6 and t[2].startswith("Sources:")][0]
check("sources cut after a whole entry", s9.endswith("2025; …"), s9[-40:])

# 10. if_empty from the table overrides the default
rules = [dict(r) for r in RULES]
for r in rules:
    if r["variable"] == "finding_02_result_achieved":
        r["if_empty"] = "firm_todo"
_, tx10 = texts(varfill.build_deck(METALCO, rules)[0])
check("table if_empty overrides the default",
      any("Eurostat reported" in t[2] and "[Firm to complete]" in t[2] for t in tx10))

# 11. a line ending in a citation keeps the citation when it has to be cut
v = dict(METALCO)
v["finding_01_context"] = "Spanish contractors moved a growing share of their steel purchasing to online channels during 2024 and 2025 across the whole of Iberia. [1, 2]"
_, tx11 = texts(varfill.build_deck(v, RULES)[0])
l11 = [t[2] for t in tx11 if t[0] == 6 and "Spanish contractors" in t[2]][0].split("\n")[0]
check("cut keeps the citation", l11.endswith("\u2026 [1, 2]") and len(l11) <= 90, l11)

# 12. slide 4 (summary), template 1.2
SUM = json.load(open(os.path.join(HERE, "summary_doc06.json")))
NOPLAN = json.load(open(os.path.join(HERE, "summary_noplan.json")))


def slide4(values):
    prs = Presentation(io.BytesIO(varfill.build_deck(values, RULES)[0]))
    s4 = list(prs.slides)[3]
    shapes = {sh.name: sh for sh in walk(s4.shapes) if getattr(sh, "has_text_frame", False)}
    col = lambda name: [p.text for p in shapes[name].text_frame.paragraphs if p.text.strip()]
    return (shapes["Title 1"].text_frame.text, col("Google Shape;881;p247"),
            col("Google Shape;882;p247"), col("Google Shape;883;p247"))


t, o, so, oc = slide4({**METALCO, **SUM})
check("slide 4: sentence title printed", t == SUM["summary_title_01"], t)
check("slide 4: four objectives", o == [SUM[f"objective_{k:02d}"] for k in range(1, 5)], o)
n_sol = sum(1 for k in range(1, 6) if SUM[f"solution_{k:02d}"])
check("slide 4: unused solution lines hidden", len(so) == n_sol, so)
check("slide 4: outcomes printed", oc[0] == SUM["outcome_01"], oc)
_, tx12 = texts(varfill.build_deck({**METALCO, **SUM}, RULES)[0])
check("slide 4: no sample text left", not any(w in t2[2] for t2 in tx12 for w in ("proven methodology", "de-risked", "Consulting firm]")))
t, o, so, oc = slide4({**METALCO, **NOPLAN})
check("slide 4, no plan: one placeholder line in Solution", so == ["[Firm to complete] The proposed approach and phases."], so)
check("slide 4, no plan: one placeholder line in Outcomes", oc == ["[Firm to complete] The expected results and deliverables."], oc)
t, o, so, oc = slide4(METALCO)
check("slide 4, nothing sent: one [Firm to complete] per column",
      o == ["[Firm to complete]"] and so == ["[Firm to complete]"] and oc == ["[Firm to complete]"] and t == "[Firm to complete]", (t, o, so, oc))

# 13. template 1.3: 1.2 + Approach section, work-plan slides drawn in place
PLAN = json.load(open(os.path.join(ROOT, "phase_architect", "runs", "final_B.json")))["plan"]
PLAN["client"] = "Atl\u00e2ntico Seguros"
RULES13 = RULES if any(r["variable"] == "summary_title_01" for r in RULES) else RULES + [{"variable": "summary_title_01", "max_chars": 90, "if_empty": "firm_todo"}]
v = dict(METALCO, section_04="Approach", summary_title_01="We will diagnose the claims delays in 12 to 16 weeks",
         objective_01="Understand why claims take too long.", solution_01="The work will run 12 to 16 weeks.",
         outcome_01="Prioritized causes of settlement delay")
b13, w13 = varfill.build_deck(v, RULES13, workplan=PLAN, version="1.3")
prs13, tx13 = texts(b13)
check("1.3: 13 slides", len(prs13.slides) == 13, len(prs13.slides))
check("1.3: no variable survives", not any(VAR_RE.search(t[2]) for t in tx13), [t for t in tx13 if VAR_RE.search(t[2])][:3])
agenda = [t[2] for t in tx13 if t[0] == 2]
check("1.3: agenda reads Context, Market Research, Approach, Next Steps",
      [a for a in agenda if a in ("Context", "Market Research", "Approach", "Next Steps")] == ["Context", "Market Research", "Approach", "Next Steps"], agenda)
check("1.3: Context slide title and columns filled",
      any(t[0] == 4 and t[2].startswith("We will diagnose") for t in tx13) and any(t[0] == 4 and "Prioritized causes" in t[2] for t in tx13))
check("1.3: empty Context lines hidden, not marked", not any(t[0] == 4 and "[Firm to complete]" in t[2] for t in tx13),
      [t[2] for t in tx13 if t[0] == 4])
check("1.3: Approach divider is 3, Next Steps divider is 4",
      any(t[0] == 7 and t[2] == "Approach" for t in tx13) and any(t[0] == 7 and t[2] == "3" for t in tx13)
      and any(t[0] == 11 and t[2] == "4" for t in tx13), [t[2] for t in tx13 if t[0] in (7, 11)])
check("1.3: work-plan slides drawn in place with titles from the plan",
      any(t[0] == 8 and t[2].startswith("Up to ") and "workstreams run in the same week" in t[2] for t in tx13)
      and any(t[0] == 9 and "deliverables" in t[2] for t in tx13)
      and any(t[0] == 10 and "steering meetings" in t[2] for t in tx13))
check("1.3: work-plan slides carry the Approach tracker", sum(1 for t in tx13 if t[0] in (8, 9, 10) and t[2] == "Approach") == 3)
_, tx13n = texts(varfill.build_deck(v, RULES13, version="1.3")[0])
check("1.3 without a plan: slides kept and marked [Firm to complete]",
      sum(1 for t in tx13n if t[0] in (8, 9, 10) and t[2].startswith("[Firm to complete]")) == 3)
try:
    varfill.build_deck(v, RULES13, version="9.9"); bad = False
except varfill.TemplateRulesError:
    bad = True
check("unknown template version is refused", bad)
check("1.2 stays the default (9 slides, no Approach)", len(texts(varfill.build_deck(METALCO, RULES)[0])[0].slides) == 9)

# 7. template 1.4 (Scenario 1): situation slide, approach overview; drop_slide (key-question slide removed Oct 18)
NAV = json.load(open(os.path.join(HERE, "metalco_1_4.json")))
tv14 = varfill.template_variables(varfill.TEMPLATES["1.4"])
check("1.4: every template variable has a table row", tv14 <= table_names, sorted(tv14 - table_names))
check("1.4: every table row has a variable in 1.4", table_names <= tv14, sorted(table_names - tv14))
pptx, warns = varfill.build_deck(NAV["values"], NAV["rules"], workplan=NAV["workplan"], version="1.4")
prs, tx = texts(pptx)
left = [t for t in tx if re.search(r"(?<![A-Za-z0-9_])(situation|belief|key_question|criteria_basis|criterion)_[0-9a-z_]+", t[2])]
n14 = len(prs.slides)  # 15, plus a steering continuation slide when a plan has more than 4 meetings
check("1.4: 15 slides + steering continuation, no warnings", n14 == 16 and not warns, (n14, warns))
check("1.4: no variable survives", not left, left[:3])
check("1.4: situation slide 6 filled", any(t[0] == 6 and t[2].startswith("Spanish contractors moved") for t in tx))
check("1.4: belief filled", any(t[0] == 6 and t[2].startswith("We believe") for t in tx))
check("1.4: no key-question slide", not any(t[2] in ("The question this proposal answers", "What a good answer must achieve") for t in tx))
check("1.4: approach overview drawn from the plan",
      any(t[0] == 9 and t[2].startswith("MetalCo decides in week") and "goes ahead" in t[2] for t in tx)
      and any(t[0] == 9 and t[2].startswith("Decision: ") for t in tx))
check("1.4: market slide keeps its own title", any(t[0] == 7 and t[2] == "Market research findings" for t in tx))
check("1.4: agenda lists the sections in variable order (Approach third, Next Steps fourth)",
      [t[2] for t in tx if t[0] == 2 and t[2] in ("Executive Summary", "Context", "Approach", "Next Steps")]
      == ["Executive Summary", "Context", "Approach", "Next Steps"], [t[2] for t in tx if t[0] == 2])
check("1.4: no SAMPLE footer; the footer variable is filled on the masters",
      "SAMPLE" not in "".join(sh.text_frame.text for m in prs.slide_masters for sh in m.shapes if sh.has_text_frame)
      and any("MetalCo | Draft for discussion" in sh.text_frame.text for m in prs.slide_masters for sh in m.shapes if sh.has_text_frame))
check("1.4: counts in words in the work-plan titles", not any(t[0] in (9, 10, 11, 12, 13) and re.search(r"\b[2-9] (phases|workstreams|deliverables|steering)", t[2]) for t in tx))
check("1.4: overview says what each phase waits for, and when the first deliverable arrives",
      any(t[0] == 9 and t[2].startswith("Starts after the ") for t in tx) and any(t[0] == 9 and "first deliverable in week" in t[2] for t in tx))
v = dict(NAV["values"]); v["timeline_title_01"] = "A partner's own timeline title"
_, txt = texts(varfill.build_deck(v, NAV["rules"], workplan=NAV["workplan"], version="1.4")[0])
check("1.4: a title variable given in the payload replaces the composed title",
      any(t[0] == 10 and t[2] == "A partner's own timeline title" for t in txt))
check("1.4: section trackers read Context on slides 6-7, Approach on 8",
      all(any(t[0] == n and t[2] == "Context" for t in tx) for n in (6, 7))
      and any(t[0] == 8 and t[2] == "Approach" for t in tx))
v = dict(NAV["values"]); v["situation_title_01"] = ""
prs2, tx2 = texts(varfill.build_deck(v, NAV["rules"], workplan=NAV["workplan"], version="1.4")[0])
check("1.4: no situation title -> slide dropped",
      len(prs2.slides) == n14 - 1 and not any(t[2].startswith("Where things stand") for t in tx2), len(prs2.slides))
# 1.5 after a dropped slide: the S3 layout adds slides; every part name stays unique (run 1619, 24 Oct)
import zipfile  # noqa: E402
b15, _w = varfill.build_deck(v, NAV["rules"], workplan=NAV["workplan"], version="1.5",
                             s3={"assumptions": [{"id": "A1", "text": "An assumption", "if_wrong": "What changes"}]})
names = zipfile.ZipFile(io.BytesIO(b15)).namelist()
p15 = Presentation(io.BytesIO(b15))
titles15 = [s.shapes.title.text_frame.text if s.shapes.title is not None else "" for s in p15.slides]
check("1.5: no situation title -> no duplicate slide parts, one assumptions slide",
      len(names) == len(set(names)) and sum(1 for t in titles15 if t.startswith("This draft rests on")) == 1,
      [n for n in names if names.count(n) > 1][:3])
prs5, tx5 = texts(varfill.build_deck(NAV["values"], NAV["rules"], workplan=None, version="1.4")[0])
check("1.4 without a plan: overview marked [Firm to complete]",
      any(t[0] == 9 and t[2].startswith("[Firm to complete] Approach overview") for t in tx5))


# Oct 24 (José): five phase-colour squares beside every timeline page, off the slide, in the colours the deck uses
def squares(b):
    prs = Presentation(io.BytesIO(b))
    return [[(x.left + x.width <= 0, str(x.fill.fore_color.rgb)) for x in sl.shapes if x.name.startswith("palette:phase")]
            for sl in prs.slides if any(x.name.startswith("palette:") for x in sl.shapes)]
sq14 = squares(pptx)
check("1.4: squares beside the timeline, off the slide, in the 1.4 colours",
      sq14 and all(len(p) == 5 and all(off for off, _ in p) for p in sq14)
      and [c for _, c in sq14[0]] == [str(c) for c in wp.PHASE_COLORS], sq14)
sq15 = squares(b15)
check("1.5: squares in the phase blues", sq15 and [c for _, c in sq15[0]] == [str(c) for c in wp.PHASE_BLUES], sq15)
check("1.3: no squares", not squares(b13), squares(b13))
# Oct 24 (José): no steering-meetings slides in 1.5; 1.4 keeps them
def steering_pages(b):
    return sum(1 for sl in Presentation(io.BytesIO(b)).slides if any(
        sh.has_text_frame and sh.text_frame.text.startswith("Steering meetings") for sh in sl.shapes))
def overview_subtitle(b):
    for sl in Presentation(io.BytesIO(b)).slides:
        if any(sh.name == "workplan:roadmap" for sh in sl.shapes):
            return [ph.text_frame.text for ph in sl.placeholders if ph.placeholder_format.idx == 1]
    return None
check("1.5: approach overview without subtitle, still in place; 1.4 keeps it",
      overview_subtitle(b15) == [] and overview_subtitle(pptx) == ["Approach overview: phases, decisions and how we work together"],
      (overview_subtitle(b15), overview_subtitle(pptx)))
check("1.5: no steering-meetings slide, 1.4 keeps it", steering_pages(b15) == 0 and steering_pages(pptx) >= 1,
      (steering_pages(b15), steering_pages(pptx)))
# Oct 25 (Valtrenix run): the deadline sentence and the deliverables slide's workstream label
dl_rel = wp._deadline_line({"text": "The CEO wants initial recommendations within 12 weeks.", "weeks_available": 12}, 12)
dl_date = wp._deadline_line({"text": "Recommendations by 1 March 2027", "weeks_available": 16}, 14)
check("deadline: one period after the brief's text; two weeks to contract named only for a date",
      ".." not in dl_rel and "contract" not in dl_rel and dl_rel.endswith("12 weeks available.")
      and dl_date.endswith("allowing two weeks to contract and start."), (dl_rel, dl_date))
lab = wp._del_label("01 Labor & agency cost baseline", "Diagnostic phase", wp.Inches(0.3),
                    wp._del_label_size(["01 Labor & agency cost baseline", "05 Scheduling and turnover"]))
check("deliverables label: a 31-character name is not clipped (8 pt for the page)",
      lab == (["01 Labor & agency cost baseline", "Diagnostic phase"], 8)
      and wp._del_label_size(["05 Scheduling and turnover"]) == 9, lab)

failed = [r for r in results if not r[1]]
for name, ok, detail in results:
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {detail}"))
print(f"\n{len(results) - len(failed)}/{len(results)} passed")
sys.exit(1 if failed else 0)
