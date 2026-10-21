"""Every deliverable reads the same wherever it appears (José, Oct 22: the numbered diamonds are the deliverables).

For each plan the work-plan slides are rendered on template 1.4 and checked deliverable by deliverable:
  - timeline: one diamond with its number, in its workstream's row, at its due week (a diamond pushed aside by a
    neighbour may sit up to one week off; such diamonds are counted as shifted);
  - deliverables page: one line 'D<n>  <name>', 'Week <due>' beside it, the meeting that reviews it, its format icon;
  - steering slides: the 'Reviews' line of the meeting that reviews it, when that meeting lists it;
  - named build (workplan.DEL_FIELDS present): the shapes carry the deliverable's variable names (deliverable_03 for
    the diamond, deliverable_03_name, _due_week, _reviewed_at, _format on its line), and
    workplan.deliverable_variables(plan) holds what the slides show.
  - value changes (named build): a new name, due week or format moves every place together; a due week outside its
    workstream or after its review meeting, a format the kind does not allow, a read-only field and an unknown
    deliverable are refused with a warning and change nothing.

Plans: the Oct 19 compiled runs, tests/metalco_1_4.json and the synthetic grid of tests/workplan_stress.py.
Usage: python3 tests/test_deliverable_vars.py [--baseline]   (--baseline: meaning checks only, no names or values)
"""
import copy, glob, io, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches
import varfill
import workplan as wp
from workplan_stress import CASES, make_plan

BASELINE = "--baseline" in sys.argv
NAMED = hasattr(wp, "DEL_FIELDS") and not BASELINE


def plans():
    out = []
    for f in sorted(glob.glob(os.path.join(ROOT, "phase_architect", "runs", "oct19_checks_formats", "*", "*", "compiled.json"))):
        out.append((os.path.relpath(f, ROOT), json.load(open(f))))
    m = json.load(open(os.path.join(HERE, "metalco_1_4.json")))
    out.append(("tests/metalco_1_4.json", m["workplan"]))
    for total, n_ws, d in CASES:
        out.append((f"stress {total}w {n_ws}ws x{d}", make_plan(total, n_ws, dels_per_ws=d)))
    return out


def render(plan, values=None):
    data, warnings = varfill.build_deck(values or {}, workplan=plan, version="1.4")
    return Presentation(io.BytesIO(data)), warnings


def text(sh):
    return sh.text_frame.text if getattr(sh, "has_text_frame", False) and sh.has_text_frame else ""


def kind_of(slide):
    t = " ".join(text(sh) for sh in slide.shapes)
    if "Plan by workstream and week" in t:
        return "timeline"
    if "Deliverables by workstream" in t:
        return "deliverables"
    if "Steering meetings –" in t:
        return "steering"
    return None


def is_diamond(sh):
    try:
        return sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.DIAMOND
    except Exception:
        return False


def check_plan(name, plan, prs, problems, stats):
    total = plan["total_duration_weeks"]["max"]
    order = wp._ws_order(plan)
    wnum = {w["id"]: i + 1 for i, w in enumerate(order)}
    dels = {d["id"]: d for d in plan["deliverables"]}
    gx = wp.L + Inches(1.85)
    col = (wp.W - Inches(1.85)) / total
    seen_dmd, seen_line = {}, {}
    reviews_text = []
    variables = wp.deliverable_variables(plan) if NAMED else None
    for si, sl in enumerate(prs.slides, 1):
        k = kind_of(sl)
        if k == "timeline":
            rows = {}
            for sh in sl.shapes:
                m = re.match(r"^(\d{2}) ", text(sh))
                if m and sh.left == wp.L and not is_diamond(sh):
                    rows[int(m.group(1))] = (sh.top, sh.top + sh.height)
            for sh in sl.shapes:
                if not (is_diamond(sh) and text(sh).strip().isdigit()):
                    continue
                did = "D" + text(sh).strip()
                seen_dmd[did] = seen_dmd.get(did, 0) + 1
                d = dels.get(did)
                if d is None:
                    problems.append(f"{name}: slide {si}: diamond {did} is not a deliverable of the plan")
                    continue
                cx, cy = sh.left + sh.width / 2, sh.top + sh.height / 2
                wk = (cx - gx) / col + 0.5
                off = abs(wk - d["due_week"])
                if off > 1.0:
                    problems.append(f"{name}: {did} diamond at week {wk:.1f}, due week {d['due_week']}")
                elif off > 0.05:
                    stats["shifted"] += 1
                r = rows.get(wnum.get(d["workstream_id"]))
                if not r or not (r[0] <= cy <= r[1]):
                    problems.append(f"{name}: {did} diamond is not in the row of {d['workstream_id']}")
                if NAMED and sh.name != wp.del_key(did):
                    problems.append(f"{name}: {did} diamond named {sh.name!r}, want {wp.del_key(did)!r}")
        elif k == "deliverables":
            by_top = {}
            for sh in sl.shapes:
                by_top.setdefault(sh.top, []).append(sh)
            for sh in sl.shapes:
                m = re.match(r"^D(\d+)\s+(.*)$", text(sh))
                if not m:
                    continue
                did = "D" + m.group(1)
                seen_line[did] = seen_line.get(did, 0) + 1
                d = dels.get(did)
                if d is None:
                    problems.append(f"{name}: slide {si}: line {did} is not a deliverable of the plan")
                    continue
                shown = m.group(2)
                if not (shown == d["name"] or (shown.endswith("…") and d["name"].startswith(shown[:-1].rstrip()))):
                    problems.append(f"{name}: {did} line reads {shown!r}, plan has {d['name']!r}")
                row = [x for x in by_top.get(sh.top, []) if x is not sh]
                due = [x for x in row if re.fullmatch(r"Week \d+", text(x))]
                if len(due) != 1 or due[0].text_frame.text != f"Week {d['due_week']}":
                    problems.append(f"{name}: {did} due reads {[text(x) for x in due]}, plan week {d['due_week']}")
                rev = [x for x in row if x.left > (due[0].left if due else sh.left + sh.width)]
                want_rev = d.get("presented_at") or "[Firm to complete]"
                if len(rev) != 1 or not (text(rev[0]) == want_rev or (text(rev[0]).endswith("…") and want_rev.startswith(text(rev[0])[:-1].rstrip()))):
                    problems.append(f"{name}: {did} reviewed-at reads {[text(x) for x in rev]}, plan {want_rev!r}")
                icons = [x for x in sl.shapes if x.shape_type == MSO_SHAPE_TYPE.PICTURE and abs(x.top + x.height / 2 - (sh.top + sh.height / 2)) < Inches(0.05) and x.left < sh.left]
                fmt = wp._format(d)
                if len(icons) != 1:
                    problems.append(f"{name}: {did} has {len(icons)} format icons")
                elif NAMED:
                    if icons[0].name != wp.del_key(did) + "_format" or icons[0]._element.nvPicPr.cNvPr.get("descr") != fmt:
                        problems.append(f"{name}: {did} icon named {icons[0].name!r} / {icons[0]._element.nvPicPr.cNvPr.get('descr')!r}, want {wp.del_key(did)}_format / {fmt}")
                elif icons[0].name != f"format:{fmt}":
                    problems.append(f"{name}: {did} icon {icons[0].name}, want format:{fmt}")
                if NAMED:
                    key = wp.del_key(did)
                    want = [(sh, key + "_name"), (due[0] if due else None, key + "_due_week"), (rev[0] if rev else None, key + "_reviewed_at")]
                    for shp, nm in want:
                        if shp is not None and shp.name != nm:
                            problems.append(f"{name}: {did} shape {text(shp)!r} named {shp.name!r}, want {nm!r}")
                    v = variables
                    if v.get(key + "_name") != d["name"] or v.get(key + "_due_week") != d["due_week"] or v.get(key + "_format") != fmt \
                            or v.get(key + "_reviewed_at") != (d.get("presented_at") or ""):
                        problems.append(f"{name}: {did} variables {[(f, v.get(key + '_' + f)) for f in wp.DEL_FIELDS]} differ from the slide")
        elif k == "steering":
            reviews_text += [text(sh) for sh in sl.shapes if re.match(r"^D\d+\s", text(sh))]
    for did in dels:
        if seen_dmd.get(did, 0) != 1:
            problems.append(f"{name}: {did} has {seen_dmd.get(did, 0)} diamonds")
        if seen_line.get(did, 0) != 1:
            problems.append(f"{name}: {did} has {seen_line.get(did, 0)} lines on the deliverables pages")
    # a 'Reviews' line names a deliverable with its plan name (the box lists up to 3 and then '+n more')
    for block in reviews_text:
        for ln in block.split("\n"):
            m = re.match(r"^D(\d+)\s+(.*)$", ln)
            if m and "D" + m.group(1) in dels:
                nm = dels["D" + m.group(1)]["name"]
                if not (m.group(2) == nm or (m.group(2).endswith("…") and nm.startswith(m.group(2)[:-1].rstrip()))):
                    problems.append(f"{name}: steering Reviews line {ln!r} differs from the plan name {nm!r}")
    stats["deliverables"] += len(dels)


def value_changes(problems):
    """Named build only: values for name, due week and format move every place; bad ones are refused."""
    plan = json.load(open(os.path.join(ROOT, "phase_architect", "runs", "oct19_checks_formats", "fmt", "n2b", "compiled.json")))
    d = plan["deliverables"][1]                                # D2
    key = wp.del_key(d["id"])
    ev = {e["label"]: e for e in plan["governance"]["events"]}
    rev_week = ev[d["presented_at"]]["week"]
    ws = {w["id"]: w for w in plan["workstreams"]}[d["workstream_id"]]
    ok_week = next((w for w in range(ws["start_week"], min(ws["end_week"], rev_week) + 1) if w != d["due_week"]), None)
    allowed = [f for f in wp.FORMAT_ALLOWED[d["kind"]] if f != wp._format(d)]
    good = {key + "_name": "Renamed deliverable for the value test", key + "_format": allowed[0]}
    if ok_week:
        good[key + "_due_week"] = str(ok_week)
    new, warns = wp.apply_deliverable_values(plan, good)
    if warns:
        problems.append(f"value test: good values gave warnings {warns}")
    nd = {x["id"]: x for x in new["deliverables"]}[d["id"]]
    if nd["name"] != good[key + "_name"] or nd["format"] != allowed[0] or (ok_week and nd["due_week"] != ok_week):
        problems.append(f"value test: values not applied: {nd}")
    if {x["id"]: x for x in plan["deliverables"]}[d["id"]]["name"] == good[key + "_name"]:
        problems.append("value test: the input plan was changed in place")
    prs, warnings = render(plan, good)
    check_plan("value test (rendered)", new, prs, problems, {"shifted": 0, "deliverables": 0})
    if not any(good[key + "_name"] in text(sh) for sl in prs.slides for sh in sl.shapes if sh.name == key + "_name"):
        problems.append("value test: the deliverables page does not show the new name")
    bad_week = rev_week + 1 if rev_week + 1 <= ws["end_week"] else ws["end_week"] + 1
    bad = {key + "_due_week": str(bad_week), key + "_format": "training" if "training" not in wp.FORMAT_ALLOWED[d["kind"]] else "nonsense",
           key + "_reviewed_at": "Some other meeting", key + "_workstream": "99 Other", "deliverable_99_name": "Nobody", key + "_name": "x" * 80}
    new2, warns2 = wp.apply_deliverable_values(plan, bad)
    nd2 = {x["id"]: x for x in new2["deliverables"]}[d["id"]]
    if nd2["due_week"] != d["due_week"] or nd2.get("format") != d.get("format") or nd2.get("presented_at") != d.get("presented_at"):
        problems.append(f"value test: a refused value changed the plan: {nd2}")
    if len(nd2["name"]) > wp.DEL_NAME_MAX:
        problems.append("value test: an 80-character name was not cut to the limit")
    for frag in ("_due_week", "_format", "_reviewed_at", "_workstream", "deliverable_99", "_name"):
        if not any(frag in w for w in warns2):
            problems.append(f"value test: no warning for {frag}: {warns2}")


def main():
    problems, stats = [], {"shifted": 0, "deliverables": 0}
    ps = plans()
    for name, plan in ps:
        prs, _ = render(plan)
        check_plan(name, plan, prs, problems, stats)
    if NAMED:
        value_changes(problems)
    mode = "baseline (meaning only)" if not NAMED else "named build (meaning, names, values)"
    print(f"{mode}: {len(ps)} plans, {stats['deliverables']} deliverables, {stats['shifted']} diamonds shifted by a neighbour")
    for p in problems[:40]:
        print("FAIL", p)
    print("all deliverables consistent" if not problems else f"{len(problems)} problems")
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
