import io, json, sys, os
sys.path.insert(0, '/home/claude/repo/jml-deck'); sys.path.insert(0, '/home/claude/repo/jml-deck/tests')
from pptx import Presentation
import workplan
from workplan_stress import check_slide
for k in sys.argv[1:]:
    plan = json.load(open(f'{k}_compiled.json'))
    data = workplan.build_workplan_deck(plan, client=plan.get('client') or 'Client')
    open(f'{k}_workplan.pptx', 'wb').write(data)
    prs = Presentation(io.BytesIO(data))
    probs = [p for i, sl in enumerate(prs.slides, 1) for p in check_slide(sl, prs.slide_width, prs.slide_height, i)]
    t = workplan.titles(plan)
    print(('FAIL ' if probs else 'PASS ') + f'{k}: {len(prs.slides)} slides | {t["timeline"]}' + (' | ' + ' | '.join(probs[:3]) if probs else ''))
