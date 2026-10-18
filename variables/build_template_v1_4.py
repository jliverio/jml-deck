"""
Builds template_v1_4.potx from template_v1_3.potx. Run from the repo root:

    python3 variables/build_template_v1_4.py

What 1.4 adds to 1.3 (Scenario 1 of the October 16 plan):

1. Situation slide, after the Context divider. Variables situation_title_01,
   situation_01..03 (where the market stands, what is changing, what it means
   for the client) and belief_01 ("We believe ..."). Written from the cited
   market findings only. situation_title_01 is drop_slide: no situation, no slide.
2. Approach overview, first slide of the Approach section. Named, not filled:
   its title placeholder is called workplan:roadmap and workplan.py draws the
   phases, the decisions that close them and a governance strip from the plan.
3. The market findings slide gets a fixed title ("Market research findings"):
   in 1.4 section_02 is "Context", so the section name no longer works as its title.
4. Oct 18 (José): section variables follow the slide order: section_03 is Approach
   (third), section_04 is Next Steps (fourth); 1.3 keeps its order. The four work-plan
   slide titles are variables (approach_title_01, timeline_title_01,
   deliverables_title_01, steering_title_01), filled by the renderer from the plan.
   The footer "JML Operations SAMPLE" on the slide masters becomes footer_01.

Section names are sent by the workflow (section_01 Executive Summary,
section_02 Context, section_03 Approach, section_04 Next Steps).

Slide order in 1.4: cover, agenda, 1 divider, summary, 2 divider, situation,
market findings, 3 divider, approach overview, timeline, deliverables,
steering meetings, 4 divider, next steps, back cover (15).

The key-question slide (governing question plus "What a good answer must
achieve") was built on Oct 17 and removed on Oct 18 at José's decision: the
summary slide already carries the objectives, solution and outcomes.
"""
import copy
import io
import os
import sys
import zipfile

from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)
import varfill  # noqa: E402
from build_template_v1_3 import duplicate_slide, set_paragraphs, add_workplan_slide  # noqa: E402

SRC = os.path.join(ROOT, "template_v1_3.potx")
DST = os.path.join(ROOT, "template_v1_4.potx")

L = 560926            # left edge of the summary slide's first column
RIGHT = 11626046      # right edge of its third column
TOP = 1715102         # column headers' top


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise SystemExit(f"shape {name!r} not found on slide")


def drop(slide, names):
    for n in names:
        sh = by_name(slide, n)
        sh._element.getparent().remove(sh._element)


def set_text(shape, text):
    """Replace a one-paragraph text box's text, keeping its first run's format."""
    p = shape.text_frame.paragraphs[0]
    runs = list(p.runs)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = text


def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h))


def tracker(slide, section):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 17:
            ph.text_frame.paragraphs[0].runs[0].text = section


def title(slide, text):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.paragraphs[0].runs[0].text = text


def clone(slide, shape):
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def no_bullets(shape, size_pt, bold=False):
    from pptx.oxml.ns import qn
    for p in shape.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        for tag in ("a:buChar", "a:buAutoNum", "a:buFont", "a:buClr", "a:buSzPct"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        if pPr.find(qn("a:buNone")) is None:
            from lxml import etree
            etree.SubElement(pPr, qn("a:buNone"))
        for r in p.runs:
            r.font.size = Pt(size_pt)
            r.font.bold = bold


def situation_slide(prs, summary):
    s = duplicate_slide(prs, summary)
    title(s, "situation_title_01")
    tracker(s, "section_02")
    drop(s, ["Google Shape;879;p247", "Google Shape;880;p247", "Google Shape;882;p247", "Google Shape;883;p247",
             "Google Shape;885;p247", "Google Shape;886;p247", "Google Shape;887;p247",
             "Google Shape;889;p247", "Google Shape;890;p247", "Google Shape;891;p247"])
    left_w = 6858000          # 7.5 in
    gap = 365760              # 0.4 in
    right_x = L + left_w + gap
    right_w = RIGHT - right_x
    head = by_name(s, "Google Shape;878;p247")
    set_text(head, "Where things stand")
    place(head, L, TOP, left_w, head.height)
    body = by_name(s, "Google Shape;881;p247")
    set_paragraphs(body, ["situation_01", "situation_02", "situation_03"])
    place(body, L, 2268393, left_w, 3766841)
    for p in body.text_frame.paragraphs:
        p.space_after = Pt(20)
        for r in p.runs:
            r.font.size = Pt(16)
    bhead = clone(s, head)
    set_text(bhead, "What we believe")
    place(bhead, right_x, TOP, right_w, head.height)
    bbody = clone(s, body)
    set_paragraphs(bbody, ["belief_01"])
    no_bullets(bbody, 20, bold=True)
    place(bbody, right_x, 2268393, right_w, 2560320)      # 2.8 in, pale panel
    from pptx.dml.color import RGBColor
    bbody.fill.solid()
    bbody.fill.fore_color.rgb = RGBColor(0xEA, 0xF4, 0xFA)
    bbody.text_frame.margin_left = bbody.text_frame.margin_right = Emu(182880)
    bbody.text_frame.margin_top = Emu(182880)
    return s


TITLE_VARS = {"roadmap": "approach_title_01", "timeline": "timeline_title_01",
              "deliverables": "deliverables_title_01", "steering": "steering_title_01"}


def swap_tokens(prs, a, b):
    """Swap two variable names everywhere on the slides (agenda, dividers, trackers, titles)."""
    import re
    pat = re.compile(r"(?<![A-Za-z0-9_])(%s|%s)(?![A-Za-z0-9_])" % (a, b))
    for sl in prs.slides:
        for sh in varfill._text_shapes(sl):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = pat.sub(lambda m: b if m.group() == a else a, r.text)


def main():
    prs = varfill._load_template(SRC)
    slides = list(prs.slides)
    if len(slides) != 13:
        raise SystemExit(f"{SRC} is not template 1.3 (expected 13 slides, found {len(slides)})")
    (cover, agenda, div1, summary, div2, market, div3, timeline, deliverables, steering,
     div4, next_steps, back) = slides
    have = {r.text for sh in summary.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    if "summary_title_01" not in have:
        raise SystemExit("slide 4 of the source is not the summary slide")

    # market findings: fixed title (the tracker keeps section_02)
    title(market, "Market research findings")

    situation = situation_slide(prs, summary)
    roadmap = add_workplan_slide(prs, timeline.slide_layout, "workplan:roadmap",
                                 "Approach overview: phases and decisions (drawn from the Phase Architect plan)",
                                 "Approach overview")

    order = [cover, agenda, div1, summary, div2, situation, market, div3, roadmap,
             timeline, deliverables, steering, div4, next_steps, back]
    lst = prs.slides._sldIdLst
    ids = list(lst)
    by_id = {s.slide_id: el for s, el in zip(prs.slides, ids)}
    order_ids = [s.slide_id for s in order]
    for el in ids:
        lst.remove(el)
    for sid in order_ids:
        lst.append(by_id[sid])

    swap_tokens(prs, "section_03", "section_04")
    # Titles inherited from template 1.1 carry a hidden bullet (a zero-width character) that pushes the text
    # right of the other titles (José, Oct 18: slides 7, 8, 13, 14). Every title gets no bullet.
    from pptx.oxml.ns import qn
    from lxml import etree
    for sl in prs.slides:
        for ph in sl.placeholders:
            if ph.placeholder_format.idx != 0:
                continue
            for p in ph.text_frame.paragraphs:
                pPr = p._p.get_or_add_pPr()
                for tag in ("a:buChar", "a:buAutoNum", "a:buFont", "a:buClr", "a:buSzPts", "a:buSzPct"):
                    for el in pPr.findall(qn(tag)):
                        pPr.remove(el)
                if pPr.find(qn("a:buNone")) is None:
                    etree.SubElement(pPr, qn("a:buNone"))
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.name.startswith("workplan:") and sh.name[9:] in TITLE_VARS:
                set_text(sh, TITLE_VARS[sh.name[9:]])
    for m in prs.slide_masters:
        for sh in m.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == "JML Operations SAMPLE":
                set_text(sh, "footer_01")
                right = sh.left + sh.width            # keep the right edge next to the page number;
                sh.width = Emu(4114800)               # 4.5 in holds "<client> | Draft for discussion"
                sh.left = Emu(right - 4114800)

    buf = io.BytesIO()
    prs.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(DST, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml",
                                    b"presentationml.template.main+xml")
            dst.writestr(item, data)
    print("wrote", DST, "slides:", len(order))


if __name__ == "__main__":
    main()
